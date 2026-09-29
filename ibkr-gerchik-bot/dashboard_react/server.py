"""Lightweight FastAPI backend for the React dashboard.

Serves memory/ data as JSON. The dashboard never connects to IBKR directly;
mutating workflows are handed off to bot jobs.
Run:  python dashboard_react/server.py
"""
from __future__ import annotations

import os
import csv
import io
import json
import math
import subprocess
import sys
import tempfile
import threading
import time
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from dataclasses import asdict, replace
from datetime import datetime, timedelta, timezone
from typing import Any, Mapping

import pandas as pd
import uvicorn
from fastapi import BackgroundTasks, FastAPI, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from zoneinfo import ZoneInfo

from dashboard import data_access as da, forecast as fc
from src.config import SETTINGS
from src.data.chart_history import daily_bars_from_intraday
from src.crypto.analysis import run_crypto_analysis
from src.crypto.config import CRYPTO_SETTINGS
from src.crypto.manual_order import execute_manual_demo_order, load_manual_order_state
from src.crypto.okx_client import OKXClient
from src.crypto.order_manager import CryptoOrderManager, CryptoOrderRequest
from src.crypto.symbols import add_crypto_symbol, default_instrument_type, normalize_okx_instrument, remove_crypto_symbol
from src.crypto.tradingview_webhook import (
    enqueue_tradingview_webhook,
    load_tradingview_state,
    process_queued_tradingview_execution,
)
from src.config import fx_pair_components
from src.forex.tradingview_webhook import handle_forex_tradingview_webhook, load_forex_tradingview_state
from src.journal.broker_reconcile import reconcile_ibkr_open_orders
from src.journal.order_journal import load_order_journal, save_review
from src.stocks.tradingview_webhook import handle_stock_tradingview_webhook, load_stock_tradingview_state
from dashboard_react.market_screener import ScreenerParams, run_market_screener
from src.scanners.inefficiency_reclaim import run_inefficiency_reclaim_screener
from src.storage.inefficiency_reclaim_store import InefficiencyReclaimStore
from src.decision.audit import DecisionAudit
from src.decision.models import AgentMode
from src.decision.provider_health import run_provider_health_check
from src.decision.strategy_control import (
    control_payload,
    create_run,
    list_runs,
    load_control,
    load_snapshot,
    manual_status_payload,
    update_control,
)
from src.decision.strategy_sources import DEFAULT_REGISTRY, refresh_current_source_snapshot
from src.decision.prompt_compiler import compile_decision_prompt
from src.decision.prompt_presets import disable_preset, duplicate_preset, get_preset, list_presets, save_preset
from src.decision.strategy_definitions import current_strategy_parameters, get_strategy_definition, list_strategy_definitions
from src.decision.strategy_sources import StrategySnapshot, snapshot_to_candidate

app = FastAPI(title="Vitaly's Trading Bot Dashboard API", docs_url=None, redoc_url=None)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["GET", "POST"],
    allow_headers=["*"],
)


@app.middleware("http")
async def disable_dashboard_cache(request: Request, call_next):
    response = await call_next(request)
    if request.url.path == "/" or request.url.path.startswith("/api/"):
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
    return response

HERE = Path(__file__).parent
app.mount("/static", StaticFiles(directory=HERE), name="static")

ET = ZoneInfo(SETTINGS.trading_hours.timezone)

_PROVIDER_TEST_LOCK = threading.Lock()
_PROVIDER_TEST_STATE: dict[str, object] = {
    "connection_status": "not_tested",
    "last_provider_test_status": "not_tested",
    "last_provider_test_at": None,
    "last_provider_latency_ms": None,
    "last_provider_error": None,
}


def _provider_test_state() -> dict[str, object]:
    with _PROVIDER_TEST_LOCK:
        return dict(_PROVIDER_TEST_STATE)


def _set_provider_test_state(**updates: object) -> dict[str, object]:
    with _PROVIDER_TEST_LOCK:
        _PROVIDER_TEST_STATE.update(updates)
        return dict(_PROVIDER_TEST_STATE)


def _decision_lab_warnings(config: object) -> list[str]:
    warnings: list[str] = []
    raw_mode = getattr(config, "mode", "off")
    mode = str(getattr(raw_mode, "value", raw_mode)).strip().lower()
    try:
        AgentMode.from_value(mode)
    except ValueError:
        warnings.append("LLM_AGENT_MODE is invalid")
    if bool(getattr(config, "allow_live_trading", False)):
        warnings.append("ALLOW_LLM_LIVE_TRADING=true")
    if not SETTINGS.paper_trading:
        warnings.append("PAPER_TRADING=false")
    if not SETTINGS.dry_run_mode:
        warnings.append("DRY_RUN_MODE=false")
    if mode == "live_autonomous":
        warnings.append("LLM_AGENT_MODE=live_autonomous")
    if mode == "ibkr_paper_autonomous":
        warnings.append("AUTONOMOUS ORDERS ENABLED / IBKR PAPER ACCOUNT ONLY")
        if not bool(getattr(config, "allow_ibkr_paper_trading", False)):
            warnings.append("ALLOW_LLM_IBKR_PAPER_TRADING=false")
        if not getattr(config, "ibkr_paper_account_allowlist", ()):
            warnings.append("LLM_IBKR_PAPER_ACCOUNT_ALLOWLIST is empty")
    return warnings


def _decision_lab_audit() -> DecisionAudit | None:
    """Open the audit DB only when it already exists; dashboard reads stay read-only."""
    path = Path(SETTINGS.decision_agent.database_path)
    if not path.exists():
        return None
    try:
        return DecisionAudit(path)
    except Exception:
        return None


def _decision_lab_positions() -> dict[str, object]:
    path = Path(SETTINGS.decision_agent.database_path).with_name("decision_lab_portfolio.json")
    if not path.exists():
        return {"owner": "LLM_AGENT", "positions": [], "closed_positions": []}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(payload, dict) or payload.get("owner") != "LLM_AGENT":
            return {"owner": "LLM_AGENT", "positions": [], "closed_positions": []}
        return {
            "owner": "LLM_AGENT",
            "equity": payload.get("equity", 0.0),
            "cash": payload.get("cash", 0.0),
            "realized_pnl": payload.get("realized_pnl", 0.0),
            "unrealized_pnl": payload.get("unrealized_pnl", 0.0),
            "daily_R": payload.get("daily_R", 0.0),
            "total_R": payload.get("total_R", 0.0),
            "positions": payload.get("positions", []),
            "closed_positions": payload.get("closed_positions", []),
        }
    except (OSError, json.JSONDecodeError):
        return {"owner": "LLM_AGENT", "positions": [], "closed_positions": [], "error": "portfolio_unreadable"}
IRS_SCHEDULED_TASKS = {
    "IBKR Bot - IRS Premarket": "Premarket",
    "IBKR Bot - IRS Hourly": "Hourly",
    "IBKR Bot - IRS Confirmation": "Confirmation",
    "IBKR Bot - IRS EOD": "EOD",
}
_IRS_TASK_CACHE: dict[str, Any] = {"loaded_at": 0.0, "tasks": None}
IRS_MIN_DAILY_HISTORY_ROWS_KEY = "IRS_MIN_DAILY_HISTORY_ROWS"
IRS_MIN_DAILY_HISTORY_ROWS_FLOOR = 20
IRS_MIN_DAILY_HISTORY_ROWS_CEILING = 5000
IRS_MINIMUM_DISPLAY_SCORE_KEY = "IRS_MINIMUM_DISPLAY_SCORE"
IRS_MINIMUM_DISPLAY_SCORE_FLOOR = 0.0


def _write_env_setting(path: Path, key: str, value: str) -> None:
    """Atomically update one allowlisted environment setting."""
    if key not in {
        IRS_MIN_DAILY_HISTORY_ROWS_KEY,
        IRS_MINIMUM_DISPLAY_SCORE_KEY,
    }:
        raise ValueError(f"Dashboard setting is not allowlisted: {key}")
    text = path.read_text(encoding="utf-8-sig") if path.exists() else ""
    newline = "\r\n" if "\r\n" in text else "\n"
    lines = text.splitlines(keepends=True)
    pattern = re.compile(rf"^(\s*{re.escape(key)}\s*=\s*)([^#\r\n]*?)(\s+#.*)?(\r?\n)?$")
    found = False
    updated_lines: list[str] = []
    for line in lines:
        match = pattern.match(line)
        if not match:
            updated_lines.append(line)
            continue
        found = True
        updated_lines.append(
            f"{match.group(1)}{value}{match.group(3) or ''}{match.group(4) or ''}"
        )
    if not found:
        if updated_lines and not updated_lines[-1].endswith(("\n", "\r")):
            updated_lines[-1] += newline
        updated_lines.append(f"{key}={value}{newline}")

    path.parent.mkdir(parents=True, exist_ok=True)
    temp_name = ""
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            newline="",
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            temp_name = handle.name
            handle.write("".join(updated_lines))
        os.replace(temp_name, path)
    finally:
        if temp_name:
            Path(temp_name).unlink(missing_ok=True)


def _parse_irs_min_daily_history_rows(value: object) -> int:
    if isinstance(value, bool):
        raise ValueError("Minimum daily history must be a whole number.")
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        raise ValueError("Minimum daily history must be a whole number.") from None
    if not math.isfinite(numeric) or not numeric.is_integer():
        raise ValueError("Minimum daily history must be a whole number.")
    rows = int(numeric)
    if not IRS_MIN_DAILY_HISTORY_ROWS_FLOOR <= rows <= IRS_MIN_DAILY_HISTORY_ROWS_CEILING:
        raise ValueError(
            "Minimum daily history must be between "
            f"{IRS_MIN_DAILY_HISTORY_ROWS_FLOOR} and {IRS_MIN_DAILY_HISTORY_ROWS_CEILING} rows."
        )
    return rows


def _set_irs_min_daily_history_rows(rows: int, *, env_path: Path | None = None) -> int:
    resolved_rows = _parse_irs_min_daily_history_rows(rows)
    updated_settings = replace(
        SETTINGS.inefficiency_reclaim,
        min_daily_history_rows=resolved_rows,
    )
    updated_settings.validate()
    _write_env_setting(
        env_path or ROOT / ".env",
        IRS_MIN_DAILY_HISTORY_ROWS_KEY,
        str(resolved_rows),
    )
    os.environ[IRS_MIN_DAILY_HISTORY_ROWS_KEY] = str(resolved_rows)
    object.__setattr__(SETTINGS, "inefficiency_reclaim", updated_settings)
    return resolved_rows


def _parse_irs_minimum_display_score(value: object) -> float:
    if isinstance(value, bool):
        raise ValueError("Inefficiency Reclaim Score must be a number.")
    try:
        score = float(value)
    except (TypeError, ValueError):
        raise ValueError("Inefficiency Reclaim Score must be a number.") from None
    maximum = float(SETTINGS.inefficiency_reclaim.minimum_order_score)
    if not math.isfinite(score) or not IRS_MINIMUM_DISPLAY_SCORE_FLOOR <= score <= maximum:
        raise ValueError(
            "Inefficiency Reclaim Score must be between "
            f"{IRS_MINIMUM_DISPLAY_SCORE_FLOOR:g} and {maximum:g}."
        )
    return score


def _set_irs_minimum_display_score(
    score: float,
    *,
    env_path: Path | None = None,
) -> float:
    resolved_score = _parse_irs_minimum_display_score(score)
    updated_settings = replace(
        SETTINGS.inefficiency_reclaim,
        minimum_display_score=resolved_score,
    )
    updated_settings.validate()
    env_value = f"{resolved_score:g}"
    _write_env_setting(
        env_path or ROOT / ".env",
        IRS_MINIMUM_DISPLAY_SCORE_KEY,
        env_value,
    )
    os.environ[IRS_MINIMUM_DISPLAY_SCORE_KEY] = env_value
    object.__setattr__(SETTINGS, "inefficiency_reclaim", updated_settings)
    return resolved_score


def _onboard_client_id(symbol: str) -> str:
    configured = os.environ.get("ONBOARD_SYMBOL_CLIENT_ID", "").strip()
    if configured:
        return configured
    seed = sum((index + 1) * ord(char) for index, char in enumerate(symbol.upper()))
    timestamp = int(datetime.now(ET).timestamp())
    return str(1000 + ((timestamp + seed) % 8000))


def _parse_stock_symbol_upload(body: dict) -> tuple[list[str], list[dict[str, str]]]:
    from src.symbol_universe import normalize_stock_symbol

    raw_values: list[str] = []
    for key in ("csv_text", "text", "symbols_text"):
        value = body.get(key)
        if isinstance(value, str) and value.strip():
            for row in csv.reader(io.StringIO(value)):
                raw_values.extend(row)
    symbols_value = body.get("symbols")
    if isinstance(symbols_value, list):
        raw_values.extend(str(item) for item in symbols_value)
    elif isinstance(symbols_value, str) and symbols_value.strip():
        for row in csv.reader(io.StringIO(symbols_value)):
            raw_values.extend(row)

    headers = {"symbol", "symbols", "ticker", "tickers", "stock", "stocks", "stock_symbol", "stock_symbols"}
    parsed: list[str] = []
    seen: set[str] = set()
    invalid: list[dict[str, str]] = []
    for raw in raw_values:
        value = str(raw or "").strip()
        if not value or value.startswith("#") or value.lower() in headers:
            continue
        try:
            symbol = normalize_stock_symbol(value)
        except ValueError as exc:
            invalid.append({"value": value, "reason": str(exc)})
            continue
        if symbol not in seen:
            seen.add(symbol)
            parsed.append(symbol)
    return parsed, invalid


def _queue_bulk_stock_onboarding(symbols: list[str]) -> dict:
    if not symbols:
        return {"pid": None, "log": None, "queue_file": None}
    SETTINGS.paths.runtime_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(ET).strftime("%Y%m%d_%H%M%S")
    queue_path = SETTINGS.paths.runtime_dir / f"bulk_onboard_symbols_{stamp}.csv"
    log_path = SETTINGS.paths.runtime_dir / f"bulk_onboard_symbols_{stamp}.log"
    with queue_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["symbol"])
        for symbol in symbols:
            writer.writerow([symbol])
    cmd = [sys.executable, "scripts/bulk_onboard_symbols.py", "--symbols-file", str(queue_path)]
    with log_path.open("a", encoding="utf-8") as log:
        log.write(f"\n--- {datetime.now(ET).isoformat()} queued {' '.join(cmd)} ---\n")
        process = subprocess.Popen(
            cmd,
            cwd=str(ROOT),
            stdout=log,
            stderr=subprocess.STDOUT,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    return {"pid": process.pid, "log": str(log_path), "queue_file": str(queue_path)}
BMSB_NEAR_CROSS_THRESHOLD_PCT = 0.75
BMSB_RECENT_CROSS_DAYS = 2
BMSB2_RECLAIM_TOLERANCE_PCT = 0.5
BMSB2_NEAR_TRIGGER_THRESHOLD_PCT = 0.75
GAUSSIAN_NEAR_THRESHOLD_PCT = 1.0
GAUSSIAN_RECENT_SIGNAL_DAYS = 2
GAUSSIAN_PERIOD = 144
GAUSSIAN_POLES = 4
GAUSSIAN_MULT = 1.414

def _safe(fn, default=None):
    try:
        return fn()
    except Exception:
        return default

def _market_session(now: datetime) -> tuple[str, str]:
    if now.weekday() >= 5:
        return "CLOSED", "Weekend"
    m = now.hour * 60 + now.minute
    if m < 240:   return "CLOSED", "Overnight"
    if m < 570:   return "PRE-MARKET", "Before 09:30"
    if m < 960:   return "OPEN", "Regular session"
    if m < 1200:  return "AFTER-HOURS", "After 16:00"
    return "CLOSED", "Overnight"


def _daily_live_frame(symbol: str, asset: str = "stock") -> pd.DataFrame:
    get_bars = da.get_crypto_bars if asset == "crypto" else da.get_bars
    daily = get_bars(symbol, "daily")
    # OKX daily candles are UTC-anchored and already include the current 24/7
    # crypto day.  The stock-only intraday stitching below applies exchange
    # session boundaries and would distort a crypto day.
    if asset == "crypto":
        return daily
    intraday = get_bars(symbol, "intraday_5m")
    if intraday.empty:
        return daily

    derived_daily = daily_bars_from_intraday(intraday)
    if derived_daily.empty:
        return daily
    if daily.empty:
        return derived_daily

    derived_dates = pd.to_datetime(derived_daily["date"], errors="coerce").dt.date
    historical_dates = set(pd.to_datetime(daily["date"], errors="coerce").dropna().dt.date)
    missing_or_live = derived_daily.loc[~derived_dates.isin(historical_dates)]
    if missing_or_live.empty:
        return daily
    return pd.concat([daily, missing_or_live], ignore_index=True).sort_values("date")


def _weekly_live_frame(symbol: str, asset: str = "stock") -> pd.DataFrame:
    """Return weekly OHLCV with current-week daily/live data stitched in."""
    daily = _daily_live_frame(symbol, asset)
    if daily.empty:
        return pd.DataFrame() if asset == "crypto" else da.get_bars(symbol, "weekly")
    frame = daily.set_index("date").sort_index()
    return (
        frame.resample("W-FRI")
        .agg({"open": "first", "high": "max", "low": "min", "close": "last", "volume": "sum"})
        .dropna(subset=["open", "high", "low", "close"])
        .reset_index()
    )


def _bmsb_scan_symbol(symbol: str, today, include_neutral: bool = False, asset: str = "stock") -> dict | None:
    weekly = _weekly_live_frame(symbol, asset)
    if weekly.empty or len(weekly) < 21:
        return None

    weekly = weekly.sort_values("date").tail(260).reset_index(drop=True)
    close = pd.to_numeric(weekly["close"], errors="coerce")
    weekly["sma20"] = close.rolling(20).mean()
    weekly["ema21"] = close.ewm(span=21, adjust=False).mean()
    ready = weekly.dropna(subset=["sma20", "ema21", "close"])
    if len(ready) < 2:
        return None

    prev = ready.iloc[-2]
    curr = ready.iloc[-1]
    prev_diff = float(prev["ema21"] - prev["sma20"])
    curr_diff = float(curr["ema21"] - curr["sma20"])
    close_price = float(curr["close"])
    if close_price <= 0:
        return None

    current_gap_pct = abs(curr_diff) / close_price * 100.0
    prev_gap_pct = abs(prev_diff) / float(prev["close"]) * 100.0 if float(prev["close"]) > 0 else None
    latest_daily = _daily_live_frame(symbol, asset)
    latest_session = None
    if not latest_daily.empty:
        latest_session = pd.to_datetime(latest_daily["date"], errors="coerce").max()
    latest_session_date = latest_session.date() if pd.notna(latest_session) else today
    recent_cutoff = today - timedelta(days=BMSB_RECENT_CROSS_DAYS)

    signal = None
    side = None
    urgency = None
    if prev_diff <= 0 < curr_diff and latest_session_date >= recent_cutoff:
        signal = "CROSSED_LONG"
        side = "LONG"
        urgency = "crossed"
    elif prev_diff >= 0 > curr_diff and latest_session_date >= recent_cutoff:
        signal = "CROSSED_EXIT"
        side = "EXIT"
        urgency = "crossed"
    elif current_gap_pct <= BMSB_NEAR_CROSS_THRESHOLD_PCT:
        signal = "NEAR_LONG" if curr_diff <= 0 else "NEAR_EXIT"
        side = "LONG" if curr_diff <= 0 else "EXIT"
        urgency = "near"

    if not signal:
        if not include_neutral:
            return None
        signal = "NO_SIGNAL"
        side = "NONE"
        urgency = "none"

    return {
        "symbol": symbol,
        "signal": signal,
        "side": side,
        "urgency": urgency,
        "last_price": round(close_price, 4),
        "ema21": round(float(curr["ema21"]), 4),
        "sma20": round(float(curr["sma20"]), 4),
        "gap": round(curr_diff, 4),
        "gap_pct": round(current_gap_pct, 4),
        "previous_gap_pct": round(prev_gap_pct, 4) if prev_gap_pct is not None else None,
        "entry_level": None,
        "exit_level": None,
        "trigger_level": None,
        "trigger_gap_pct": round(current_gap_pct, 4),
        "cross_date": str(latest_session_date),
        "weekly_bar": str(pd.to_datetime(curr["date"]).date()),
        "threshold_pct": BMSB_NEAR_CROSS_THRESHOLD_PCT,
    }


def _bmsb_strategy2_scan_symbol(symbol: str, today, include_neutral: bool = False, asset: str = "stock") -> dict | None:
    weekly = _weekly_live_frame(symbol, asset)
    if weekly.empty or len(weekly) < 21:
        return None

    weekly = weekly.sort_values("date").tail(260).reset_index(drop=True)
    close = pd.to_numeric(weekly["close"], errors="coerce")
    weekly["sma20"] = close.rolling(20).mean()
    weekly["ema21"] = close.ewm(span=21, adjust=False).mean()
    ready = weekly.dropna(subset=["sma20", "ema21", "close"])
    if len(ready) < 2:
        return None

    tolerance_mult = BMSB2_RECLAIM_TOLERANCE_PCT / 100.0
    ready = ready.copy()
    ready["band_upper"] = ready[["sma20", "ema21"]].max(axis=1)
    ready["band_lower"] = ready[["sma20", "ema21"]].min(axis=1)
    ready["entry_level"] = ready["band_upper"] * (1.0 - tolerance_mult)
    ready["exit_level"] = ready["band_lower"] * (1.0 + tolerance_mult)

    prev = ready.iloc[-2]
    curr = ready.iloc[-1]
    close_price = float(curr["close"])
    if close_price <= 0:
        return None

    latest_daily = _daily_live_frame(symbol, asset)
    latest_session = None
    if not latest_daily.empty:
        latest_session = pd.to_datetime(latest_daily["date"], errors="coerce").max()
    latest_session_date = latest_session.date() if pd.notna(latest_session) else today
    recent_cutoff = today - timedelta(days=BMSB_RECENT_CROSS_DAYS)

    entry_level = float(curr["entry_level"])
    exit_level = float(curr["exit_level"])
    prev_entry_level = float(prev["entry_level"])
    prev_exit_level = float(prev["exit_level"])
    prev_close = float(prev["close"])

    long_entry = (
        prev_close <= prev_entry_level
        and close_price > entry_level
        and latest_session_date >= recent_cutoff
    )
    long_exit = (
        prev_close >= prev_exit_level
        and close_price < exit_level
        and latest_session_date >= recent_cutoff
    )

    signal = None
    side = None
    urgency = None
    trigger_level = None
    trigger_gap_pct = None

    if long_entry:
        signal = "BMSB2_LONG_ENTRY"
        side = "LONG"
        urgency = "crossed"
        trigger_level = entry_level
        trigger_gap_pct = 0.0
    elif long_exit:
        signal = "BMSB2_EXIT"
        side = "EXIT"
        urgency = "crossed"
        trigger_level = exit_level
        trigger_gap_pct = 0.0
    else:
        neutral_candidates = [
            ("BMSB2_NEAR_ENTRY", "LONG", entry_level, abs(entry_level - close_price) / close_price * 100.0),
            ("BMSB2_NEAR_EXIT", "EXIT", exit_level, abs(close_price - exit_level) / close_price * 100.0),
        ]
        candidates = []
        if close_price <= entry_level:
            distance = (entry_level - close_price) / close_price * 100.0
            candidates.append(("BMSB2_NEAR_ENTRY", "LONG", entry_level, distance))
        if close_price >= exit_level:
            distance = (close_price - exit_level) / close_price * 100.0
            candidates.append(("BMSB2_NEAR_EXIT", "EXIT", exit_level, distance))
        candidates = [item for item in candidates if item[3] <= BMSB2_NEAR_TRIGGER_THRESHOLD_PCT]
        if candidates:
            signal, side, trigger_level, trigger_gap_pct = min(candidates, key=lambda item: item[3])
            urgency = "near"

    ema21 = float(curr["ema21"])
    sma20 = float(curr["sma20"])
    band_gap_pct = abs(ema21 - sma20) / close_price * 100.0

    if not signal:
        if not include_neutral:
            return None
        _, nearest_side, nearest_level, nearest_gap_pct = min(neutral_candidates, key=lambda item: item[3])
        signal = "NO_SIGNAL"
        side = "NONE"
        urgency = "none"
        trigger_level = nearest_level
        trigger_gap_pct = nearest_gap_pct

    return {
        "symbol": symbol,
        "signal": signal,
        "side": side,
        "urgency": urgency,
        "last_price": round(close_price, 4),
        "ema21": round(ema21, 4),
        "sma20": round(sma20, 4),
        "gap": round(ema21 - sma20, 4),
        "gap_pct": round(band_gap_pct, 4),
        "entry_level": round(entry_level, 4),
        "exit_level": round(exit_level, 4),
        "trigger_level": round(float(trigger_level), 4) if trigger_level is not None else None,
        "trigger_gap_pct": round(float(trigger_gap_pct), 4) if trigger_gap_pct is not None else None,
        "nearest_side": nearest_side if signal == "NO_SIGNAL" else side,
        "cross_date": str(latest_session_date),
        "weekly_bar": str(pd.to_datetime(curr["date"]).date()),
        "threshold_pct": BMSB2_NEAR_TRIGGER_THRESHOLD_PCT,
        "tolerance_pct": BMSB2_RECLAIM_TOLERANCE_PCT,
    }


def _gaussian_alpha(period: int, poles: int) -> float:
    beta = (1 - math.cos(2 * math.pi / period)) / (math.sqrt(2) ** (2 / poles) - 1)
    return -beta + math.sqrt(beta * beta + 2 * beta)


def _gaussian_filter(values: list[float], period: int = GAUSSIAN_PERIOD, poles: int = GAUSSIAN_POLES) -> list[float]:
    alpha = _gaussian_alpha(period, poles)
    f1 = f2 = f3 = f4 = None
    out: list[float] = []
    for value in values:
        x = float(value) if math.isfinite(float(value)) else 0.0
        if f1 is None:
            f1 = f2 = f3 = f4 = x
        else:
            f1 = alpha * x + (1 - alpha) * f1
            f2 = alpha * f1 + (1 - alpha) * f2
            f3 = alpha * f2 + (1 - alpha) * f3
            f4 = alpha * f3 + (1 - alpha) * f4
        out.append(f1 if poles == 1 else f2 if poles == 2 else f3 if poles == 3 else f4)
    return out


def _cross_over(prev_a: float | None, prev_b: float | None, a: float, b: float) -> bool:
    return (
        prev_a is not None
        and prev_b is not None
        and math.isfinite(prev_a)
        and math.isfinite(prev_b)
        and math.isfinite(a)
        and math.isfinite(b)
        and prev_a <= prev_b
        and a > b
    )


def _cross_under(prev_a: float | None, prev_b: float | None, a: float, b: float) -> bool:
    return (
        prev_a is not None
        and prev_b is not None
        and math.isfinite(prev_a)
        and math.isfinite(prev_b)
        and math.isfinite(a)
        and math.isfinite(b)
        and prev_a >= prev_b
        and a < b
    )


def _gaussian_scan_symbol(symbol: str, today, include_neutral: bool = False, asset: str = "stock") -> dict | None:
    daily = _daily_live_frame(symbol, asset)
    if daily.empty or len(daily) < 3:
        return None

    frame = daily.sort_values("date").tail(420).reset_index(drop=True).copy()
    for column in ("open", "high", "low", "close"):
        frame[column] = pd.to_numeric(frame[column], errors="coerce")
    frame = frame.dropna(subset=["date", "high", "low", "close"]).reset_index(drop=True)
    if len(frame) < 3:
        return None

    src = ((frame["high"] + frame["low"] + frame["close"]) / 3).astype(float).tolist()
    ranges: list[float] = []
    closes = frame["close"].astype(float).tolist()
    highs = frame["high"].astype(float).tolist()
    lows = frame["low"].astype(float).tolist()
    for idx, high in enumerate(highs):
        prev_close = closes[idx - 1] if idx > 0 else closes[idx]
        ranges.append(max(high - lows[idx], abs(high - prev_close), abs(lows[idx] - prev_close)))

    mid_values = _gaussian_filter(src)
    range_values = _gaussian_filter(ranges)
    position = 0
    latest_signal: tuple[str, str, str, str] | None = None
    latest_daily_date = pd.to_datetime(frame["date"].iloc[-1], errors="coerce").date()
    recent_cutoff = today - timedelta(days=GAUSSIAN_RECENT_SIGNAL_DAYS)

    for idx, close in enumerate(closes):
        mid = mid_values[idx]
        width = GAUSSIAN_MULT * range_values[idx]
        upper = mid + width
        lower = mid - width
        prev_mid = mid_values[idx - 1] if idx > 0 else None
        prev_close = closes[idx - 1] if idx > 0 else None
        prev_upper = mid_values[idx - 1] + GAUSSIAN_MULT * range_values[idx - 1] if idx > 0 else None
        prev_lower = mid_values[idx - 1] - GAUSSIAN_MULT * range_values[idx - 1] if idx > 0 else None
        bullish = prev_mid is not None and mid > prev_mid

        signal = None
        side = None
        if bullish and _cross_over(prev_close, prev_upper, close, upper) and position <= 0:
            signal = "GAUSSIAN_LONG_ENTRY"
            side = "LONG"
            position = 1
        elif position > 0 and _cross_under(prev_close, prev_upper, close, upper):
            signal = "GAUSSIAN_LONG_EXIT"
            side = "EXIT"
            position = 0
        if signal:
            latest_signal = (signal, side or "NONE", "crossed", str(pd.to_datetime(frame["date"].iloc[idx]).date()))

    close_price = float(closes[-1])
    mid = float(mid_values[-1])
    upper = mid + GAUSSIAN_MULT * float(range_values[-1])
    lower = mid - GAUSSIAN_MULT * float(range_values[-1])
    prev_mid = float(mid_values[-2])
    bullish = mid > prev_mid
    near_long_gap = abs(upper - close_price) / close_price * 100.0 if close_price > 0 else None
    exit_gap = abs(close_price - upper) / close_price * 100.0 if close_price > 0 else None

    signal = side = urgency = None
    signal_date = None
    if latest_signal:
        signal, side, urgency, signal_date = latest_signal
        try:
            if pd.to_datetime(signal_date).date() < recent_cutoff:
                signal = side = urgency = signal_date = None
        except Exception:
            signal = side = urgency = signal_date = None
    if not signal and bullish and near_long_gap is not None and close_price <= upper and near_long_gap <= GAUSSIAN_NEAR_THRESHOLD_PCT:
        signal = "GAUSSIAN_NEAR_LONG"
        side = "LONG"
        urgency = "near"
        signal_date = str(latest_daily_date)
    elif not signal and position > 0 and exit_gap is not None and close_price >= upper and exit_gap <= GAUSSIAN_NEAR_THRESHOLD_PCT:
        signal = "GAUSSIAN_NEAR_EXIT"
        side = "EXIT"
        urgency = "near"
        signal_date = str(latest_daily_date)

    if not signal:
        if not include_neutral:
            return None
        signal = "NO_SIGNAL"
        side = "NONE"
        urgency = "none"
        signal_date = str(latest_daily_date)

    gap = close_price - upper
    gap_pct = abs(gap) / close_price * 100.0 if close_price > 0 else None
    return {
        "symbol": symbol,
        "signal": signal,
        "side": side,
        "urgency": urgency,
        "last_price": round(close_price, 4),
        "ema21": round(upper, 4),
        "sma20": round(mid, 4),
        "gap": round(gap, 4),
        "gap_pct": round(gap_pct, 4) if gap_pct is not None else None,
        "previous_gap_pct": None,
        "entry_level": round(upper, 4),
        "exit_level": round(upper, 4),
        "trigger_level": round(upper, 4),
        "trigger_gap_pct": round(gap_pct, 4) if gap_pct is not None else None,
        "cross_date": signal_date,
        "daily_bar": str(latest_daily_date),
        "weekly_bar": str(latest_daily_date),
        "threshold_pct": GAUSSIAN_NEAR_THRESHOLD_PCT,
        "channel_mid": round(mid, 4),
        "channel_upper": round(upper, 4),
        "channel_lower": round(lower, 4),
    }


@app.get("/")
def index():
    return FileResponse(HERE / "index.html")


@app.get("/api/meta")
def api_meta():
    now = datetime.now(ET)
    session, session_detail = _market_session(now)
    health = _safe(da.source_health, [])
    return {
        "clock": now.strftime("%H:%M:%S"),
        "session": session,
        "session_detail": session_detail,
        "paper": SETTINGS.paper_trading,
        "dry_run": SETTINGS.dry_run_mode,
        "health": [asdict(h) | {"path": str(h.path), "updated_at": h.updated_at.isoformat() if h.updated_at else None} for h in health],
    }


@app.get("/api/decision-lab/status")
def api_decision_lab_status():
    config = SETTINGS.decision_agent
    raw_mode = getattr(config, "mode", "off")
    mode = str(getattr(raw_mode, "value", raw_mode)).strip().lower()
    audit = _decision_lab_audit()
    status = audit.status() if audit is not None else {
        "candidates": 0, "decisions": 0, "model_decisions": 0,
        "risk_decisions": 0, "execution_links": 0, "position_events": 0,
        "outcomes": 0, "shadow_decisions": 0,
        "internal_paper_executions": 0, "ibkr_paper_executions": 0,
        "live_executions": 0, "executions_by_mode": {},
    }
    status.pop("database", None)
    diagnostic_state = _provider_test_state()
    _, autonomous_worker = _autonomous_stock_worker_status()
    return {
        "mode": mode,
        "provider": str(config.provider),
        "model": str(config.model or config.local_model or ""),
        "news_enabled": bool(config.use_news),
        "multi_provider_shadow": bool(config.multi_provider_shadow),
        "decision_workers": int(config.decision_workers),
        "workers": int(config.decision_workers),
        "queue_depth": int(config.decision_queue_depth),
        "candidate_expiry": float(config.candidate_expiry_seconds),
        "paper_trading": bool(SETTINGS.paper_trading),
        "dry_run": bool(SETTINGS.dry_run_mode),
        "allow_llm_live_trading": bool(config.allow_live_trading),
        "warnings": _decision_lab_warnings(config),
        **diagnostic_state,
        "database_available": audit is not None,
        "paper_portfolio": _decision_lab_positions(),
        "ibkr_paper_autonomous_enabled": mode == "ibkr_paper_autonomous",
        "allow_llm_ibkr_paper_trading": bool(getattr(config, "allow_ibkr_paper_trading", False)),
        # The dashboard does not connect to IBKR merely to render status.  A
        # verified account must therefore be reported as not checked, never
        # inferred from an allowlist or a socket port.
        "paper_account_verified": None if mode == "ibkr_paper_autonomous" else False,
        "paper_account_allowlisted": None if mode == "ibkr_paper_autonomous" else False,
        "paper_account_status": autonomous_worker.get("account_status", "UNKNOWN"),
        "worker_running": bool(autonomous_worker.get("process_alive") and autonomous_worker.get("status") in {"running", "blocked"}),
        "worker_state": autonomous_worker.get("state", "STOPPED"),
        "worker_last_heartbeat": autonomous_worker.get("last_heartbeat"),
        "broker_status": autonomous_worker.get("broker_status", "UNKNOWN"),
        "account_status": autonomous_worker.get("account_status", "UNKNOWN"),
        "allowlist_status": autonomous_worker.get("allowlist_status", "UNKNOWN"),
        "provider_status": autonomous_worker.get("provider_status", "UNKNOWN"),
        "provider_last_checked_at": autonomous_worker.get("provider_last_checked_at"),
        "provider_latency_ms": autonomous_worker.get("provider_latency_ms"),
        "provider_last_error": autonomous_worker.get("provider_last_error"),
        "current_session": autonomous_worker.get("current_session", "MARKET CLOSED"),
        "autonomous_entry_enabled": autonomous_worker.get("autonomous_entry_enabled", False),
        "blocked_reason": autonomous_worker.get("blocked_reason", "WORKER_HEARTBEAT_UNAVAILABLE"),
        "last_preflight_at": autonomous_worker.get("last_preflight_at"),
        "autonomous_stock_worker": autonomous_worker,
        **manual_status_payload(),
        **status,
    }


@app.get("/api/decision-lab/decisions")
def api_decision_lab_decisions(request: Request):
    audit = _decision_lab_audit()
    if audit is None:
        return {"decisions": []}
    query = request.query_params
    decisions = audit.list_decisions(
        symbol=query.get("symbol"),
        action=query.get("action"),
        provider=query.get("provider"),
        mode=query.get("mode"),
        limit=int(query.get("limit", "100")),
    )
    return {"decisions": _decision_table_projection(decisions)}


def _decision_table_projection(decisions: list[dict[str, object]]) -> list[dict[str, object]]:
    """Keep list/run responses lightweight; full prompt bodies require a detail request."""
    for row in decisions:
        try:
            candidate = json.loads(str(row.get("candidate_json") or "{}"))
        except (TypeError, json.JSONDecodeError):
            candidate = {}
        metadata = candidate.get("metadata") if isinstance(candidate, dict) else {}
        metadata = metadata if isinstance(metadata, dict) else {}
        row["source"] = metadata.get("strategy_source") or "UNKNOWN"
        row["source_signal"] = metadata.get("source_signal") or "UNKNOWN"
        row["candidate_class"] = metadata.get("candidate_class") or "UNKNOWN"
        row.pop("candidate_json", None)
        row.pop("risk_reasons", None)
        row.pop("order_ids_json", None)
    return decisions


@app.get("/api/decision-lab/positions")
def api_decision_lab_positions():
    return _decision_lab_positions()


@app.get("/api/decision-lab/decision/{decision_id}")
def api_decision_lab_decision(decision_id: str):
    """Return one sanitized inspector record; this is the only prompt-detail read path."""
    audit = _decision_lab_audit()
    if audit is None:
        raise HTTPException(503, "decision audit unavailable")
    decision = audit.get_decision(decision_id)
    if decision is None:
        raise HTTPException(404, "decision not found")
    return {"decision": decision}


@app.get("/api/decision-lab/run/{run_id}")
def api_decision_lab_run_detail(run_id: str):
    """Return run counters plus lightweight decision summaries for the inspector."""
    audit = _decision_lab_audit()
    if audit is None:
        raise HTTPException(503, "decision audit unavailable")
    run = next((item for item in list_runs(limit=200) if str(item.get("run_id")) == str(run_id)), None)
    if run is None:
        raise HTTPException(404, "run not found")
    decisions = _decision_table_projection(audit.list_decisions_for_run(run_id))
    return {"run": run, "decisions": decisions, "decision_count": len(decisions)}


@app.get("/api/decision-lab/performance")
def api_decision_lab_performance():
    audit = _decision_lab_audit()
    return audit.performance() if audit is not None else {"total_outcomes": 0, "closed_outcomes": 0, "pnl": 0.0, "r": 0.0}


@app.get("/api/decision-lab/providers")
def api_decision_lab_providers():
    audit = _decision_lab_audit()
    return {"providers": audit.provider_health() if audit is not None else []}


@app.get("/api/decision-lab/strategies")
def api_decision_lab_strategies():
    """Return the registry-backed source catalog; no market or broker call."""
    return {"strategies": DEFAULT_REGISTRY.infos()}


@app.get("/api/decision-lab/strategy-definitions")
def api_decision_lab_strategy_definitions(request: Request):
    source = str(request.query_params.get("source") or "").strip().lower()
    if source:
        definition = get_strategy_definition(source)
        return {"definitions": [definition.to_dict()], "parameters": current_strategy_parameters(source)}
    return {"definitions": list_strategy_definitions()}


@app.get("/api/decision-lab/prompt-presets")
def api_decision_lab_prompt_presets(request: Request):
    source = request.query_params.get("strategy_source")
    return {"presets": list_presets(source)}


@app.post("/api/decision-lab/prompt-presets/duplicate")
async def api_decision_lab_duplicate_preset(request: Request):
    body = await request.json()
    if not isinstance(body, dict):
        raise HTTPException(400, "request body must be an object")
    try:
        return {"preset": duplicate_preset(str(body.get("prompt_id") or ""), name=str(body.get("name") or ""))}
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


@app.post("/api/decision-lab/prompt-presets")
async def api_decision_lab_save_preset(request: Request):
    body = await request.json()
    if not isinstance(body, dict):
        raise HTTPException(400, "request body must be an object")
    try:
        return {"preset": save_preset(
            strategy_source=str(body.get("strategy_source") or ""),
            name=str(body.get("name") or ""),
            prompt_text=str(body.get("prompt_text") or ""),
            description=str(body.get("description") or ""),
            parent_prompt_id=str(body.get("parent_prompt_id") or "") or None,
            as_new_version=bool(body.get("as_new_version", False)),
        )}
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


@app.post("/api/decision-lab/prompt-presets/{prompt_id}/disable")
def api_decision_lab_disable_preset(prompt_id: str):
    try:
        return {"preset": disable_preset(prompt_id)}
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


def _preview_candidate(source: str) -> tuple[object, tuple[str, ...]]:
    snapshot = load_snapshot()
    sources = snapshot.get("sources", {}) if isinstance(snapshot, dict) else {}
    selected: Mapping[str, object] | None = None
    if source == "all":
        for items in sources.values() if isinstance(sources, dict) else ():
            if isinstance(items, list) and items:
                selected = items[0]
                break
    elif isinstance(sources, dict) and isinstance(sources.get(source), list) and sources[source]:
        selected = sources[source][0]
    if not isinstance(selected, Mapping):
        raise ValueError("no persisted candidate is available for preview")
    snapshot_obj = StrategySnapshot(
        scan_id=str(selected.get("scan_id") or snapshot.get("scan_id") or "preview"),
        source=str(selected.get("source") or source),
        symbol=str(selected.get("symbol") or "").upper(),
        strategy=str(selected.get("strategy") or ""),
        signal_state=str(selected.get("signal_state") or ""),
        candidate_class=str(selected.get("candidate_class") or "WATCH_CANDIDATE"),
        price=selected.get("price"), entry=selected.get("entry"), stop=selected.get("stop"), target=selected.get("target"),
        reward_risk=selected.get("reward_risk"), atr=selected.get("atr"), score=selected.get("score"),
        source_timestamp=selected.get("source_timestamp"), bar_timestamp=selected.get("bar_timestamp"),
        direction=str(selected.get("direction") or "none"), metadata=selected.get("metadata") if isinstance(selected.get("metadata"), Mapping) else {},
    )
    return snapshot_to_candidate(snapshot_obj)


@app.post("/api/decision-lab/prompt-preview")
async def api_decision_lab_prompt_preview(request: Request):
    body = await request.json()
    if not isinstance(body, dict):
        raise HTTPException(400, "request body must be an object")
    source = str(body.get("strategy_source") or "").strip().lower()
    try:
        candidate, allowed_actions = _preview_candidate(source)
        preset = get_preset(str(body.get("prompt_id"))) if body.get("prompt_id") else None
        compiled = compile_decision_prompt(
            strategy_source=source,
            candidate=candidate,
            position_context=_decision_lab_positions().get("positions", []),
            allowed_actions=allowed_actions,
            prompt_preset=preset,
        )
        config = SETTINGS.decision_agent
        user_prompt = json.dumps(compiled.user_payload, separators=(",", ":"), ensure_ascii=False)
        return {
            "preview": compiled.to_dict(),
            "provider_request": {
                "model": str(config.model or config.local_model or ""),
                "temperature": 0,
                "max_tokens": int(getattr(config, "max_completion_tokens", 900)),
                "response_format": {"type": "json_object"},
                "messages": [
                    {"role": "system", "content": compiled.system_prompt},
                    {"role": "user", "content": user_prompt},
                ],
            },
            "provider_called": False,
        }
    except (ValueError, KeyError) as exc:
        raise HTTPException(409, str(exc)) from exc


@app.get("/api/decision-lab/control")
def api_decision_lab_control():
    """Return control plus the worker's persisted source snapshot.

    The status/read path stays fast and read-only. The persistent worker is
    responsible for refreshing current source data; the MANUAL run endpoint
    performs its required current-data validation before queueing a run.
    """
    return control_payload(registry=DEFAULT_REGISTRY)


@app.post("/api/decision-lab/control")
async def api_decision_lab_update_control(request: Request):
    """Persist only the UI control selection; execution remains worker-owned."""
    try:
        body = await request.json()
    except Exception as exc:
        raise HTTPException(400, "request body must be JSON") from exc
    if not isinstance(body, dict):
        raise HTTPException(400, "request body must be an object")
    try:
        return update_control(
            analysis_mode=body.get("analysis_mode"),
            strategy_source=body.get("strategy_source"),
            prompt_preset_id=body.get("prompt_preset_id"),
            updated_by="decision_lab",
            registry=DEFAULT_REGISTRY,
        )
    except (ValueError, KeyError) as exc:
        raise HTTPException(400, str(exc)) from exc


@app.get("/api/decision-lab/runs")
def api_decision_lab_runs(request: Request):
    try:
        limit = int(request.query_params.get("limit", "20"))
    except ValueError as exc:
        raise HTTPException(400, "limit must be an integer") from exc
    return {"runs": list_runs(limit=limit)}


@app.post("/api/decision-lab/run")
async def api_decision_lab_run(request: Request):
    """Queue one bounded MANUAL run for the persistent worker.

    This endpoint first materializes the current persisted strategy data into
    an immutable analysis snapshot, then writes a durable request/run record.
    It never creates a market-data or IBKR client and never executes a broker
    operation.
    """
    try:
        body = await request.json()
    except Exception as exc:
        raise HTTPException(400, "request body must be JSON") from exc
    if body is None:
        body = {}
    if not isinstance(body, dict):
        raise HTTPException(400, "request body must be an object")
    control = load_control(registry=DEFAULT_REGISTRY)
    if control.get("analysis_mode") != "manual":
        raise HTTPException(409, "RUN ANALYSIS requires MANUAL mode")
    source = str(body.get("strategy_source") or control.get("strategy_source") or "").strip().lower()
    try:
        DEFAULT_REGISTRY.validate_selection(source)
        refresh_current_source_snapshot(registry=DEFAULT_REGISTRY, force=True)
        selected_snapshot = control_payload(
            registry=DEFAULT_REGISTRY,
            strategy_source=source,
        ).get("selected_source_snapshot") or {}
        if not isinstance(selected_snapshot, dict) or not selected_snapshot.get("available"):
            raise HTTPException(409, {"error": "SOURCE_DATA_UNAVAILABLE", "strategy_source": source})
        latest_scan_id = str(selected_snapshot.get("scan_id") or "").strip()
        snapshot_timestamp = str(selected_snapshot.get("snapshot_timestamp") or "").strip()
        if not latest_scan_id or not snapshot_timestamp:
            raise HTTPException(409, {"error": "SOURCE_DATA_UNAVAILABLE", "strategy_source": source})
        run = create_run(
            analysis_mode="manual",
            strategy_source=source,
            prompt_preset_id=str(control.get("prompt_preset_id") or "") or None,
            scan_id=latest_scan_id,
            snapshot_timestamp=snapshot_timestamp,
        )
    except RuntimeError as exc:
        raise HTTPException(409, str(exc)) from exc
    except HTTPException:
        raise
    except (ValueError, KeyError) as exc:
        raise HTTPException(400, str(exc)) from exc
    return {
        "run": run,
        "run_id": run["run_id"],
        "state": run["status"],
        "strategy_source": run["strategy_selection"],
        "scan_id": run["scan_id"],
        "snapshot_timestamp": run["snapshot_timestamp"],
        "analysis_snapshot_id": run.get("analysis_snapshot_id", run["scan_id"]),
        "source_timestamp": selected_snapshot.get("source_timestamp"),
    }


@app.post("/api/decision-lab/test-provider")
def api_decision_lab_test_provider():
    """Run the shared configured-provider WAIT diagnostic without trading state."""
    config = SETTINGS.decision_agent
    result = run_provider_health_check(config)
    connection_status = "connected" if result.connected else "error"
    state = _set_provider_test_state(
        connection_status=connection_status,
        last_provider_test_status=connection_status,
        last_provider_test_at=result.checked_at,
        last_provider_latency_ms=result.latency_ms,
        last_provider_error=result.error,
    )
    payload = {
        "ok": result.connected,
        "provider": result.provider,
        "model": result.model,
        **state,
    }
    if result.connected:
        payload["message"] = "Provider connectivity test successful."
    else:
        payload["error"] = result.error or "UNKNOWN_PROVIDER_ERROR"
    return payload


@app.post("/api/system/hard-reset")
def api_system_hard_reset():
    reset_script = ROOT / "run_dashboard_hard_reset.cmd"
    if not reset_script.exists():
        raise HTTPException(500, f"Reset script not found: {reset_script}")
    try:
        flags = getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0) | getattr(subprocess, "CREATE_NO_WINDOW", 0)
        subprocess.Popen(
            ["cmd", "/c", str(reset_script)],
            cwd=str(ROOT),
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            stdin=subprocess.DEVNULL,
            creationflags=flags,
        )
        return {
            "ok": True,
            "message": "Hard reset started. Dashboard services will close and reopen in a few seconds.",
        }
    except Exception as exc:
        raise HTTPException(500, str(exc))


@app.post("/api/system/logout")
def api_system_logout():
    """Start the fixed dashboard-service shutdown script for the UI logout action."""
    stop_script = ROOT / "stop_react_dashboard.cmd"
    if not stop_script.exists():
        raise HTTPException(500, f"Stop script not found: {stop_script}")
    try:
        flags = getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0) | getattr(subprocess, "CREATE_NO_WINDOW", 0)
        subprocess.Popen(
            ["cmd", "/c", str(stop_script), "--no-pause"],
            cwd=str(ROOT),
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            stdin=subprocess.DEVNULL,
            creationflags=flags,
        )
        return {
            "ok": True,
            "message": "Dashboard shutdown started. You can close this tab.",
        }
    except Exception as exc:
        raise HTTPException(500, str(exc))


@app.post("/api/system/fetch-ibkr-candles")
def api_system_fetch_ibkr_candles():
    """Queue a one-shot IBKR candle refresh, ignoring normal market-hours gating."""
    try:
        SETTINGS.paths.runtime_dir.mkdir(parents=True, exist_ok=True)
        log_path = SETTINGS.paths.runtime_dir / "manual_market_data_fetch.log"
        client_id = os.environ.get("MANUAL_MARKET_DATA_CLIENT_ID", "117").strip() or "117"
        cmd = [
            sys.executable,
            "-m",
            "src.main",
            "--job",
            "market_data",
            "--once",
            "--force",
            "--interval-seconds",
            "300",
            "--client-id",
            str(client_id),
        ]
        with log_path.open("a", encoding="utf-8") as log:
            log.write(f"\n--- {datetime.now(ET).isoformat()} queued {' '.join(cmd)} ---\n")
            process = subprocess.Popen(
                cmd,
                cwd=str(ROOT),
                stdout=log,
                stderr=subprocess.STDOUT,
                stdin=subprocess.DEVNULL,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
        return {
            "ok": True,
            "pid": process.pid,
            "client_id": client_id,
            "log": str(log_path),
            "message": "Queued manual IBKR candle refresh. It will fetch latest watchlist bars even outside the normal market-data window.",
        }
    except Exception as exc:
        raise HTTPException(500, str(exc))


def _pid_alive(pid: object) -> bool:
    try:
        pid_int = int(pid)
        if pid_int <= 0:
            return False
        if os.name == "nt":
            import ctypes

            kernel32 = ctypes.windll.kernel32
            handle = kernel32.OpenProcess(0x1000, False, pid_int)  # PROCESS_QUERY_LIMITED_INFORMATION
            if not handle:
                return False
            try:
                exit_code = ctypes.c_ulong()
                if not kernel32.GetExitCodeProcess(handle, ctypes.byref(exit_code)):
                    return False
                return exit_code.value == 259  # STILL_ACTIVE
            finally:
                kernel32.CloseHandle(handle)
        os.kill(pid_int, 0)
        return True
    except PermissionError:
        return True
    except Exception:
        return False


def _read_json_file(path: Path) -> dict | None:
    try:
        if not path.exists():
            return None
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        return data if isinstance(data, dict) else None
    except Exception:
        return None


def _iso_age_seconds(value: object) -> float | None:
    if not value:
        return None
    try:
        text = str(value).replace("Z", "+00:00")
        parsed = datetime.fromisoformat(text)
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return max(0.0, (datetime.now(timezone.utc) - parsed.astimezone(timezone.utc)).total_seconds())
    except Exception:
        return None


def _service_status(label: str, ok: bool, *, age: float | None = None, detail: str = "", status: str | None = None, icon: str = "settings_ethernet") -> dict:
    if status is None:
        status = "online" if ok else "offline"
    return {
        "label": label,
        "status": status,
        "ok": bool(ok),
        "age_seconds": None if age is None else round(float(age), 1),
        "detail": detail,
        "icon": icon,
    }


def _query_irs_scheduled_tasks() -> list[dict[str, Any]] | None:
    cached_tasks = _IRS_TASK_CACHE.get("tasks")
    if time.monotonic() - float(_IRS_TASK_CACHE.get("loaded_at") or 0.0) < 15:
        return cached_tasks
    if os.name != "nt":
        _IRS_TASK_CACHE.update({"loaded_at": time.monotonic(), "tasks": None})
        return None

    names = ", ".join(f"'{name.replace(chr(39), chr(39) * 2)}'" for name in IRS_SCHEDULED_TASKS)
    command = (
        f"$names=@({names}); $rows=@(); "
        "foreach($name in $names){"
        "$task=Get-ScheduledTask -TaskName $name -ErrorAction SilentlyContinue; "
        "if($null -ne $task){"
        "$info=Get-ScheduledTaskInfo -TaskName $name -ErrorAction SilentlyContinue; "
        "$rows += [pscustomobject]@{"
        "task_name=$name; state=$task.State.ToString(); "
        "last_run=if($info -and $info.LastRunTime -gt [datetime]::MinValue){$info.LastRunTime.ToString('o')}else{$null}; "
        "next_run=if($info -and $info.NextRunTime -gt [datetime]::MinValue){$info.NextRunTime.ToString('o')}else{$null}; "
        "last_result=if($info){$info.LastTaskResult}else{$null}"
        "}}}; ConvertTo-Json -InputObject @($rows) -Compress"
    )
    try:
        completed = subprocess.run(
            ["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", command],
            capture_output=True,
            text=True,
            timeout=8,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            check=True,
        )
        parsed = json.loads(completed.stdout.strip() or "[]")
        tasks = parsed if isinstance(parsed, list) else [parsed]
        tasks = [task for task in tasks if isinstance(task, dict)]
    except Exception:
        tasks = None
    if not tasks:
        manifest = _read_json_file(SETTINGS.paths.runtime_dir / "irs_scheduler_tasks.json")
        manifest_tasks = (manifest or {}).get("tasks")
        if isinstance(manifest_tasks, list):
            tasks = [
                {
                    "task_name": str(task.get("task_name") or ""),
                    "state": "Ready",
                    "next_run": _next_manifest_run(task.get("times"), now=datetime.now(ET)),
                    "source": "installation_manifest",
                }
                for task in manifest_tasks
                if isinstance(task, dict) and task.get("task_name")
            ]
    _IRS_TASK_CACHE.update({"loaded_at": time.monotonic(), "tasks": tasks})
    return tasks


def _next_manifest_run(times: object, *, now: datetime) -> str | None:
    if not isinstance(times, list):
        return None
    parsed_times: list[tuple[int, int]] = []
    for value in times:
        try:
            hour_text, minute_text = str(value).split(":", 1)
            parsed_times.append((int(hour_text), int(minute_text)))
        except (TypeError, ValueError):
            continue
    for day_offset in range(8):
        day = now + timedelta(days=day_offset)
        if day.weekday() >= 5:
            continue
        for hour, minute in sorted(parsed_times):
            candidate = day.replace(hour=hour, minute=minute, second=0, microsecond=0)
            if candidate >= now:
                return candidate.isoformat()
    return None


def _scheduled_time(value: object) -> datetime | None:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return parsed.replace(tzinfo=ET) if parsed.tzinfo is None else parsed.astimezone(ET)
    except (TypeError, ValueError):
        return None


def _irs_schedule_status(
    *,
    tasks: list[dict[str, Any]] | None = None,
    runtime: dict[str, Any] | None = None,
    now: datetime | None = None,
) -> dict:
    scheduled_tasks = _query_irs_scheduled_tasks() if tasks is None else tasks
    runtime_state = (
        _read_json_file(SETTINGS.paths.runtime_dir / "irs_scheduler_status.json")
        if runtime is None
        else runtime
    ) or {}
    current = now or datetime.now(ET)
    if scheduled_tasks is None:
        return _service_status(
            "IRS Scheduler",
            False,
            detail="Windows Task Scheduler status is unavailable",
            status="unknown",
            icon="schedule",
        )

    by_name = {
        str(task.get("task_name") or task.get("TaskName") or ""): task
        for task in scheduled_tasks
    }
    installed = [name for name in IRS_SCHEDULED_TASKS if name in by_name]
    missing = [IRS_SCHEDULED_TASKS[name] for name in IRS_SCHEDULED_TASKS if name not in by_name]
    all_installed = len(installed) == len(IRS_SCHEDULED_TASKS)

    upcoming: list[tuple[datetime, str]] = []
    task_running = False
    for name, task in by_name.items():
        task_running = task_running or str(task.get("state") or task.get("State") or "").lower() == "running"
        next_run = _scheduled_time(task.get("next_run") or task.get("NextRunTime"))
        if name in IRS_SCHEDULED_TASKS and next_run and next_run >= current:
            upcoming.append((next_run, IRS_SCHEDULED_TASKS[name]))
    upcoming.sort(key=lambda item: item[0])
    next_detail = (
        f"next {upcoming[0][1]} {upcoming[0][0].strftime('%a %H:%M')}"
        if upcoming
        else "no next run reported"
    )

    runtime_status = str(runtime_state.get("status") or "").lower()
    runtime_pid = runtime_state.get("pid")
    runtime_alive = _pid_alive(runtime_pid) if runtime_pid else False
    runtime_active = runtime_status in {"starting", "running"} and runtime_alive
    mode = str(runtime_state.get("mode") or "IRS").replace("_", " ").title()
    age = _iso_age_seconds(runtime_state.get("updated_at"))

    if runtime_active or task_running:
        connected = runtime_active and bool(runtime_state.get("connected"))
        status = "connected" if connected else "running"
        detail = (
            f"{mode} - IBKR connected - pid {runtime_pid}"
            if connected
            else f"{mode} - connecting to IBKR"
            if runtime_active
            else "Windows task is running - connection status pending"
        )
        return _service_status(
            "IRS Scheduler",
            True,
            age=age,
            detail=f"{detail}; {next_detail}",
            status=status,
            icon="schedule",
        )

    if not all_installed:
        detail = f"{len(installed)}/4 tasks installed"
        if missing:
            detail += f"; missing {', '.join(missing)}"
        return _service_status(
            "IRS Scheduler",
            False,
            age=age,
            detail=detail,
            status="partial" if installed else "not installed",
            icon="schedule",
        )

    if runtime_status == "failed":
        error = str(runtime_state.get("error") or "last IRS run failed")
        return _service_status(
            "IRS Scheduler",
            False,
            age=age,
            detail=f"{error}; {next_detail}",
            status="failed",
            icon="schedule",
        )

    if runtime_status == "complete":
        finished_age = _iso_age_seconds(runtime_state.get("finished_at"))
        finished_text = (
            f"finished {int(finished_age)}s ago"
            if finished_age is not None
            else "last run finished"
        )
        return _service_status(
            "IRS Scheduler",
            True,
            age=finished_age,
            detail=f"{mode} {finished_text}; {next_detail}",
            status="finished",
            icon="schedule",
        )

    return _service_status(
        "IRS Scheduler",
        True,
        age=age,
        detail=f"4 tasks ready; {next_detail}",
        status="scheduled",
        icon="schedule",
    )


def _lock_status(path: Path, *, stale_after_seconds: float) -> tuple[bool, float | None, str]:
    if not path.exists():
        return False, None, "lock not present"
    try:
        age = max(0.0, time.time() - path.stat().st_mtime)
        text = path.read_text(encoding="utf-8", errors="ignore").strip()
        pid = text.split("|", 1)[0].strip() if text else ""
        alive = _pid_alive(pid) if pid else age <= stale_after_seconds
        ok = alive and age <= stale_after_seconds
        detail = f"pid {pid} · {int(age)}s ago" if pid else f"{int(age)}s ago"
        if not ok and age > stale_after_seconds:
            detail = f"stale · {detail}"
        detail = detail.replace("\ufffd", "-")
        return ok, age, detail
    except Exception as exc:
        return False, None, f"lock read error: {exc}"


def _market_collector_window_status(now: datetime | None = None) -> str:
    current = now or datetime.now(ET)
    if current.weekday() >= 5:
        return "waiting"
    start = current.replace(
        hour=SETTINGS.trading_hours.market_open_hour,
        minute=SETTINGS.trading_hours.market_open_minute,
        second=0,
        microsecond=0,
    )
    end = current.replace(
        hour=SETTINGS.trading_hours.market_data_collector_end_hour,
        minute=SETTINGS.trading_hours.market_data_collector_end_minute,
        second=0,
        microsecond=0,
    )
    return "online" if start <= current <= end else "waiting"


def _latest_job_log_status(job_name: str) -> dict:
    runtime_dir = SETTINGS.paths.runtime_dir
    newest: dict = {"status": "unknown", "ok": False, "age": None, "detail": "no job log found"}
    try:
        logs = sorted(runtime_dir.glob("application.*.log"), key=lambda path: path.stat().st_mtime, reverse=True)
    except Exception:
        return newest
    marker = f"job={job_name}"
    for path in logs:
        try:
            text = path.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            continue
        if marker not in text:
            continue
        stat = path.stat()
        age = max(0.0, time.time() - stat.st_mtime)
        pid_match = re.search(rf"job={re.escape(job_name)}\b[^\n]*\bpid=(\d+)", text)
        pid = pid_match.group(1) if pid_match else ""
        completed = (
            "Job completed:" in text
            and (
                f"'job': '{job_name}'" in text
                or f'"job": "{job_name}"' in text
                or f'"job":"{job_name}"' in text
            )
        )
        if completed:
            return {
                "status": "complete",
                "ok": age < 24 * 60 * 60,
                "age": age,
                "detail": f"completed {int(age)}s ago" + (f" - pid {pid}" if pid else ""),
                "pid": pid,
                "path": str(path),
            }
        alive = _pid_alive(pid) if pid else False
        if alive:
            return {
                "status": "running",
                "ok": True,
                "age": age,
                "detail": f"pid {pid} - log {int(age)}s ago",
                "pid": pid,
                "path": str(path),
            }
        newest = {
            "status": "complete" if age < 24 * 60 * 60 else "stale",
            "ok": age < 24 * 60 * 60,
            "age": age,
            "detail": f"last log {int(age)}s ago" + (f" - pid {pid} ended" if pid else ""),
            "pid": pid,
            "path": str(path),
        }
        break
    return newest


def _premarket_status() -> dict:
    status = _latest_job_log_status("premarket")
    if status.get("status") == "running":
        return status

    today = datetime.now(ET).strftime("%Y%m%d")
    report = SETTINGS.paths.reports_dir / f"premarket_levels_{today}.xlsx"
    if report.exists():
        age = max(0.0, time.time() - report.stat().st_mtime)
        return {
            "status": "complete",
            "ok": True,
            "age": age,
            "detail": f"today's report ready - {int(age)}s ago",
        }

    current = datetime.now(ET)
    if current.weekday() < 5 and current.hour < SETTINGS.trading_hours.market_open_hour:
        return {
            "status": "pending",
            "ok": False,
            "age": status.get("age"),
            "detail": "no completed pre-open report yet",
        }
    return status


def _autonomous_stock_worker_status() -> tuple[dict, dict]:
    """Return read-only worker service and sanitized lifecycle evidence.

    The dashboard reads this credential-free heartbeat only.  It never creates
    an IBKR client and never infers broker identity from configuration or port.
    """
    unknown = {
        "service": "autonomous_stock_worker",
        "state": "STOPPED",
        "status": "offline",
        "pid": None,
        "heartbeat_age_seconds": None,
        "last_heartbeat": None,
        "current_session": "MARKET CLOSED",
        "llm_mode": "off",
        "provider_status": "UNKNOWN",
        "provider_last_checked_at": None,
        "provider_latency_ms": None,
        "provider_last_error": None,
        "broker_status": "UNKNOWN",
        "account_status": "UNKNOWN",
        "allowlist_status": "UNKNOWN",
        "broker_positions_readable": False,
        "broker_open_orders_readable": False,
        "broker_executions_readable": False,
        "autonomous_entry_enabled": False,
        "blocked_reason": "WORKER_HEARTBEAT_UNAVAILABLE",
        "last_preflight_at": None,
        "last_error": None,
        "process_alive": False,
    }
    paths = getattr(SETTINGS, "paths", None)
    if paths is None or not getattr(paths, "runtime_dir", None):
        service = _service_status(
            "Autonomous Stock Worker",
            False,
            detail="worker heartbeat unavailable in this runtime context",
            status="offline",
            icon="bolt",
        )
        return service, unknown

    heartbeat_path = paths.runtime_dir / "autonomous_stock_worker.json"
    heartbeat = _read_json_file(heartbeat_path) or {}
    if not heartbeat:
        service = _service_status(
            "Autonomous Stock Worker",
            False,
            detail="worker heartbeat unavailable",
            status="offline",
            icon="bolt",
        )
        return service, unknown

    age = _iso_age_seconds(heartbeat.get("last_heartbeat"))
    pid = heartbeat.get("pid")
    process_alive = _pid_alive(pid)
    fresh = age is not None and age <= 20.0
    raw_state = str(heartbeat.get("state") or "STOPPED").upper()
    if raw_state in {"STOPPED", "STOPPING"} and not process_alive:
        service_state = "stopped"
    elif raw_state == "BLOCKED" and process_alive and fresh:
        service_state = "blocked"
    elif process_alive and fresh:
        # MARKET_CLOSED, READY, and DEGRADED all describe a resident worker;
        # the safety fields below distinguish whether entries are allowed.
        service_state = "running"
    else:
        service_state = "stale"
    worker_ok = process_alive and fresh and service_state in {"running", "blocked"}

    session = str(heartbeat.get("current_session") or "MARKET CLOSED").upper()
    mode = str(heartbeat.get("llm_mode") or "off")
    provider = str(heartbeat.get("provider_status") or "UNKNOWN").upper()
    broker = str(heartbeat.get("broker_status") or "UNKNOWN").upper()
    account = str(heartbeat.get("account_status") or "UNKNOWN").upper()
    allowlist = str(heartbeat.get("allowlist_status") or "UNKNOWN").upper()
    entries = bool(heartbeat.get("autonomous_entry_enabled", False))
    blocked_reason = str(heartbeat.get("blocked_reason") or "").upper()

    # A stopped worker cannot provide current safety evidence.  A stale worker
    # keeps its last categorical evidence only so the UI can mark it STALE.
    if service_state in {"stopped", "offline"}:
        provider = broker = account = allowlist = "UNKNOWN"
        entries = False
        blocked_reason = "WORKER_STOPPED"
    elif not blocked_reason:
        blocked_reason = "" if entries else "PREFLIGHT_BLOCKED"

    entry_text = "new entries enabled" if entries else "new entries blocked"
    detail = f"{session} - mode {mode} - broker {broker} - provider {provider} - {entry_text}"
    if blocked_reason:
        detail += f" - {blocked_reason}"
    service = _service_status(
        "Autonomous Stock Worker",
        worker_ok,
        age=age,
        detail=detail,
        status=service_state,
        icon="bolt",
    )
    evidence = {
        "service": "autonomous_stock_worker",
        "state": raw_state if service_state != "stale" else "STALE",
        "status": service_state,
        "worker_running": worker_ok,
        "pid": pid,
        "heartbeat_age_seconds": None if age is None else round(float(age), 1),
        "last_heartbeat": heartbeat.get("last_heartbeat"),
        "current_session": session,
        "llm_mode": mode,
        "provider_status": provider,
        "provider_last_checked_at": heartbeat.get("provider_last_checked_at"),
        "provider_latency_ms": heartbeat.get("provider_latency_ms"),
        "provider_last_error": str(heartbeat.get("provider_last_error") or "") or None,
        "broker_status": broker,
        "account_status": account,
        "allowlist_status": allowlist,
        "broker_positions_readable": bool(heartbeat.get("broker_positions_readable", False)),
        "broker_open_orders_readable": bool(heartbeat.get("broker_open_orders_readable", False)),
        "broker_executions_readable": bool(heartbeat.get("broker_executions_readable", False)),
        "autonomous_entry_enabled": entries,
        "blocked_reason": blocked_reason,
        "last_preflight_at": heartbeat.get("last_preflight_at"),
        "last_error": "worker_error" if heartbeat.get("last_error") else None,
        "process_alive": process_alive,
    }
    return service, evidence


@app.get("/api/services")
def api_services():
    from src.execution import order_requests as oq

    execute_age = _safe(oq.worker_age_seconds, None)
    execute_ok = bool(_safe(oq.worker_is_alive, False))
    execute_detail = (
        f"heartbeat {int(execute_age)}s ago"
        if execute_age is not None
        else "no heartbeat"
    )
    premarket = _premarket_status()
    irs_scheduler = _irs_schedule_status()
    autonomous_worker, autonomous_evidence = _autonomous_stock_worker_status()

    market_lock = SETTINGS.paths.runtime_dir / "market_data_collector.lock"
    _, market_age, market_lock_detail = _lock_status(market_lock, stale_after_seconds=900)
    market_pid = None
    if market_lock.exists():
        try:
            market_pid = market_lock.read_text(encoding="utf-8", errors="ignore").strip().split("|", 1)[0].strip()
        except Exception:
            market_pid = None
    market_process_alive = _pid_alive(market_pid) if market_pid else False
    market_window_status = _market_collector_window_status()
    market_status = market_window_status if market_process_alive else "offline"
    market_ok = market_process_alive
    market_alive_detail = (
        f"pid {market_pid} - startup lock {int(market_age)}s old"
        if market_pid and market_age is not None
        else "collector process alive"
    )
    market_detail = (
        f"waiting for 09:30 ET - {market_alive_detail}"
        if market_process_alive and market_window_status == "waiting"
        else f"collector alive - {market_alive_detail}"
        if market_process_alive
        else market_lock_detail
    )

    crypto_heartbeat_path = CRYPTO_SETTINGS.memory_dir / "runtime" / "worker.json"
    crypto_hb = _read_json_file(crypto_heartbeat_path)
    crypto_age = _iso_age_seconds((crypto_hb or {}).get("updated_at"))
    crypto_interval = float((crypto_hb or {}).get("interval_seconds") or CRYPTO_SETTINGS.collect_interval_seconds or 300)
    crypto_stale_after = max(crypto_interval * 3, 900.0)
    crypto_pid = (crypto_hb or {}).get("pid")
    crypto_process_alive = _pid_alive(crypto_pid) if crypto_pid else False
    crypto_fresh = crypto_age is not None and crypto_age <= crypto_stale_after
    crypto_raw_status = str((crypto_hb or {}).get("status") or "unknown").lower()
    crypto_status = (
        crypto_raw_status
        if bool(crypto_hb) and crypto_process_alive and crypto_fresh
        else "stale"
        if bool(crypto_hb) and crypto_process_alive
        else "offline"
    )
    crypto_ok = crypto_status in {"ok", "sleeping", "running"}
    crypto_detail = (
        f"{(crypto_hb or {}).get('status', 'unknown')} · pid {crypto_pid} · {int(crypto_age)}s ago"
        if crypto_age is not None
        else "no heartbeat"
    )

    crypto_detail = crypto_detail.replace("\ufffd", "-")

    ibkr_ok = execute_ok or market_ok
    ibkr_detail = (
        "via execute worker / market-data collector"
        if execute_ok and market_ok
        else "via execute worker"
        if execute_ok
        else "via market-data collector"
        if market_ok
        else "no live IBKR-connected service detected"
    )

    services = [
        _service_status("Dashboard API", True, age=0, detail="FastAPI responding", icon="dashboard"),
        autonomous_worker,
        _service_status(
            "Pre-Open",
            bool(premarket.get("ok")),
            age=premarket.get("age"),
            detail=str(premarket.get("detail") or ""),
            status=str(premarket.get("status") or "unknown"),
            icon="wb_twilight",
        ),
        _service_status("Market Data", market_ok, age=market_age, detail=market_detail, status=market_status, icon="database"),
        _service_status("Crypto Worker", crypto_ok, age=crypto_age, detail=crypto_detail, status=crypto_status, icon="currency_bitcoin"),
        irs_scheduler,
        _service_status("IBKR Bridge", ibkr_ok, age=min([a for a in [execute_age, market_age] if a is not None], default=None), detail=ibkr_detail, icon="account_balance"),
    ]
    return {
        "updated_at": datetime.now(ET).isoformat(),
        "services": services,
        "autonomous_stock_worker": autonomous_evidence,
        "all_ok": all(service["ok"] for service in services),
    }


@app.get("/api/dashboard")
def api_dashboard():
    watchlist = _safe(da.load_watchlist, {})
    decisions = _safe(da.load_daily_decisions, {})
    attempts = _safe(lambda: da.all_decision_attempts(decisions), [])
    positions = _safe(da.load_tracked_positions, [])
    snapshot = _safe(da.load_intraday_snapshot, {})
    news_items = _safe(lambda: da.blocked_news_summary(watchlist), [])

    from collections import Counter
    signals = Counter(str(a.get("signal", "NONE")).upper() for a in attempts)
    buys, sells = signals["BUY"], signals["SELL"]
    tradable = sum(1 for p in watchlist.values() if isinstance(p, dict) and p.get("levels") and not p.get("news_blocked"))
    blocked = sum(1 for p in watchlist.values() if isinstance(p, dict) and p.get("news_blocked"))
    executed = snapshot.get("executed", []) if isinstance(snapshot, dict) else []
    skipped  = snapshot.get("skipped", []) if isinstance(snapshot, dict) else []
    fill_total = len(executed) + len(skipped)
    fill_rate = len(executed) / fill_total if fill_total else 0.0

    recent_attempts = sorted(attempts, key=lambda a: str(a.get("timestamp", "")), reverse=True)[:14]
    opp_rows = []
    for sym, plan in sorted((watchlist or {}).items()):
        if not isinstance(plan, dict): continue
        levels = plan.get("levels", [])
        spacing = plan.get("level_spacing", {}) or {}
        blocked_sym = bool(plan.get("news_blocked"))
        status = "BLOCKED" if blocked_sym else ("READY" if levels else "NO LEVELS")
        opp_rows.append({
            "sym": sym, "status": status, "trade_levels": len(levels),
            "room": spacing.get("reason", "-"),
            "daily_atr": plan.get("daily_atr"),
            "price": spacing.get("current_price"),
            "strategy": (levels[0].get("strategy") if levels else "-"),
        })

    return {
        "metrics": {
            "watchlist": len(watchlist),
            "trade_ready": tradable,
            "news_blocked": blocked,
            "attempts": len(attempts),
            "action_signals": buys + sells,
            "buys": buys, "sells": sells,
            "fill_rate": fill_rate,
            "fill_executed": len(executed),
            "fill_total": fill_total,
        },
        "news_items": news_items,
        "symbols": sorted((watchlist or {}).keys()),
        "opp_rows": opp_rows[:25],
        "attempt_log": recent_attempts,
        "risk": {
            "risk_per_trade": SETTINGS.risk.risk_per_trade,
            "max_daily_loss_pct": SETTINGS.risk.max_daily_loss_pct,
            "min_reward_risk_ratio": SETTINGS.risk.min_reward_risk_ratio,
            "max_positions": SETTINGS.risk.max_positions,
            "max_spread_pct": SETTINGS.risk.max_spread_pct,
            "max_open_risk_pct": SETTINGS.risk.max_open_risk_pct,
        },
    }


@app.get("/api/watchlist")
def api_watchlist():
    from src.symbol_universe import load_stock_symbols

    watchlist = _safe(da.load_watchlist, {})
    configured_symbols = _safe(load_stock_symbols, [])
    rows = []

    def _num(value):
        try:
            result = float(value)
            return result if pd.notna(result) else None
        except (TypeError, ValueError):
            return None

    symbols = sorted(set(configured_symbols or []) | set((watchlist or {}).keys()))
    for symbol in symbols:
        plan = (watchlist or {}).get(symbol, {})
        has_plan = isinstance(plan, dict) and bool(plan)
        if not isinstance(plan, dict):
            plan = {}

        spacing = plan.get("level_spacing", {}) or {}
        levels = plan.get("levels", []) or []
        blocked = bool(plan.get("news_blocked"))
        status = "PENDING" if not has_plan else ("BLOCKED" if blocked else ("READY" if levels else "NO LEVELS"))

        daily = _safe(lambda s=symbol: _daily_live_frame(s), pd.DataFrame())
        last_price = _num(spacing.get("current_price"))
        prev_close = None
        change = None
        change_pct = None
        day_open = day_high = day_low = volume = None
        last_bar = None

        if daily is not None and not daily.empty:
            daily = daily.sort_values("date").dropna(subset=["close"])
            if not daily.empty:
                last = daily.iloc[-1]
                last_price = _num(last.get("close")) or last_price
                day_open = _num(last.get("open"))
                day_high = _num(last.get("high"))
                day_low = _num(last.get("low"))
                volume = _num(last.get("volume"))
                last_bar = str(last.get("date")) if last.get("date") is not None else None
                if len(daily) >= 2:
                    prev_close = _num(daily.iloc[-2].get("close"))
                    if last_price is not None and prev_close not in (None, 0):
                        change = last_price - prev_close
                        change_pct = change / prev_close * 100.0

        rows.append({
            "symbol": symbol,
            "status": status,
            "last_price": round(last_price, 4) if last_price is not None else None,
            "prev_close": round(prev_close, 4) if prev_close is not None else None,
            "change": round(change, 4) if change is not None else None,
            "change_pct": round(change_pct, 4) if change_pct is not None else None,
            "day_open": round(day_open, 4) if day_open is not None else None,
            "day_high": round(day_high, 4) if day_high is not None else None,
            "day_low": round(day_low, 4) if day_low is not None else None,
            "volume": int(volume) if volume is not None else None,
            "daily_atr": round(_num(plan.get("daily_atr")), 4) if _num(plan.get("daily_atr")) is not None else None,
            "technical_atr": round(_num(plan.get("technical_atr")), 4) if _num(plan.get("technical_atr")) is not None else None,
            "trade_levels": len(levels),
            "raw_levels": len(plan.get("raw_levels", []) or []),
            "news": "BLOCKED" if blocked else "CLEAR",
            "last_bar": last_bar,
            "security_type": plan.get("security_type", "-"),
            "room": spacing.get("reason", "-"),
        })

    gainers = sum(1 for row in rows if (row.get("change") or 0) > 0)
    losers = sum(1 for row in rows if (row.get("change") or 0) < 0)
    unchanged = len(rows) - gainers - losers
    ready = sum(1 for row in rows if row.get("status") == "READY")
    pending = sum(1 for row in rows if row.get("status") == "PENDING")

    return {
        "rows": rows,
        "metrics": {
            "symbols": len(rows),
            "ready": ready,
            "pending": pending,
            "gainers": gainers,
            "losers": losers,
            "unchanged": unchanged,
        },
        "updated_at": datetime.now(ET).isoformat(),
    }


def _screener_frame(symbol: str, timeframe: str, asset: str = "stock") -> pd.DataFrame:
    selected_asset = "crypto" if str(asset or "stock").strip().lower() == "crypto" else "stock"
    get_bars = da.get_crypto_bars if selected_asset == "crypto" else da.get_bars
    timeframe = str(timeframe or "1D").upper()
    if timeframe == "1D":
        return _daily_live_frame(symbol, selected_asset)
    if timeframe == "4H":
        stored = get_bars(symbol, "intraday_4h")
        if not stored.empty:
            return stored
        intraday = get_bars(symbol, "intraday_5m")
        if intraday.empty:
            return intraday
        frame = intraday.set_index("date").sort_index()
        return (
            frame.resample("240min", origin="start_day", offset="0min" if selected_asset == "crypto" else "30min")
            .agg({"open": "first", "high": "max", "low": "min", "close": "last", "volume": "sum"})
            .dropna(subset=["open", "high", "low", "close"])
            .reset_index()
        )
    if timeframe == "1H":
        stored = get_bars(symbol, "intraday_1h")
        if not stored.empty:
            return stored
        intraday = get_bars(symbol, "intraday_5m")
        if intraday.empty:
            return intraday
        frame = intraday.set_index("date").sort_index()
        return (
            frame.resample("60min", origin="start_day", offset="0min" if selected_asset == "crypto" else "30min")
            .agg({"open": "first", "high": "max", "low": "min", "close": "last", "volume": "sum"})
            .dropna(subset=["open", "high", "low", "close"])
            .reset_index()
        )
    return _daily_live_frame(symbol, selected_asset)


@app.get("/api/market-screener")
@app.get("/api/crypto/market-screener")
def api_market_screener(
    request: Request,
    timeframe: str = "1D",
    strategies: str = "LP1,LP2,PRB1,PRB2",
    side: str = "ALL",
    min_score: float = 0.0,
    rr: float = 2.0,
    risk_pct: float = 0.5,
    equity: float = 100000.0,
    anchor_date: str = "",
    entry_atr_pct: float = 0.50,
    level_tolerance_atr: float = 0.10,
    open_tolerance_atr: float = 0.15,
    lp1_delta_break: float = 0.05,
    lp1_delta_close: float = 0.03,
    lp2_delta_break: float = 0.05,
    lp2_delta_close: float = 0.03,
    prb1_delta_break: float = 0.05,
    prb1_delta_close: float = 0.05,
    prb2_delta_break: float = 0.05,
    prb2_delta_close: float = 0.05,
    lp1_volume_mult: float = 1.0,
    lp2_volume_mult: float = 1.0,
    prb_volume_mult: float = 1.2,
    range_cap: float = 2.5,
    two_bar_range_cap: float = 4.0,
    stop_buffer_atr: float = 0.25,
    stop_variant: str = "strategy_default",
    approach_lookback: int = 4,
    approach_min_atr: float = 0.05,
    gap_max_atr: float = 1.2,
    chase_max_atr: float = 0.30,
    lp2_overextended_atr: float = 1.2,
    slippage_per_share: float = 0.01,
    fees_per_share: float = 0.005,
    require_complete_quality: bool = False,
):
    from src.symbol_universe import load_stock_symbols

    selected_asset = "crypto" if request.url.path.startswith("/api/crypto/") else "stock"
    if selected_asset == "crypto":
        crypto_state = _safe(da.load_crypto_dashboard_state, {})
        watchlist = crypto_state.get("watchlist", {}) if isinstance(crypto_state, dict) else {}
        configured = _safe(da.load_crypto_symbols, [])
        configured_symbols = [
            str(row.get("inst_id") or "").upper()
            for row in configured
            if isinstance(row, dict) and row.get("inst_id")
        ]
        if CRYPTO_SETTINGS.symbols_file.exists():
            symbols = sorted(set(configured_symbols))
        else:
            symbols = sorted(set(configured_symbols) | set((_safe(da.crypto_bars_index, {}) or {}).keys()))
    else:
        watchlist = _safe(da.load_watchlist, {})
        configured_symbols = _safe(load_stock_symbols, [])
        symbols = sorted(set(configured_symbols or []) | set((watchlist or {}).keys()))
    selected_strategies = tuple(
        strategy
        for strategy in (item.strip().upper() for item in str(strategies or "").split(","))
        if strategy in {"LP1", "LP2", "PRB1", "PRB2"}
    ) or ("LP1", "LP2", "PRB1", "PRB2")
    selected_stop_variant = str(stop_variant or "strategy_default").strip().lower()
    allowed_stop_variants = {
        "strategy_default",
        "behind_level",
        "behind_signal_bar",
        "behind_two_bar_structure",
        "behind_day1",
    }
    if selected_stop_variant not in allowed_stop_variants:
        selected_stop_variant = "strategy_default"
    params = ScreenerParams(
        asset_class=selected_asset,
        timeframe=str(timeframe or "1D").upper(),
        anchor_date=str(anchor_date or "").strip() or None,
        strategies=selected_strategies,
        side_filter=str(side or "ALL").upper(),
        score_min=max(0.0, min(1.0, float(min_score or 0.0))),
        rr=max(1.0, min(3.0, float(rr or 2.0))),
        risk_pct=max(0.0001, min(0.05, float(risk_pct or 0.5) / 100.0)),
        equity=max(1.0, float(equity or 100000.0)),
        entry_atr_pct=max(0.20, min(0.80, float(entry_atr_pct or 0.50))),
        level_tolerance_atr=max(0.05, min(0.20, float(level_tolerance_atr or 0.10))),
        open_tolerance_atr=max(0.05, min(0.30, float(open_tolerance_atr or 0.15))),
        lp1_delta_break=max(0.01, min(0.30, float(lp1_delta_break or 0.05))),
        lp1_delta_close=max(0.01, min(0.20, float(lp1_delta_close or 0.03))),
        lp2_delta_break=max(0.01, min(0.30, float(lp2_delta_break or 0.05))),
        lp2_delta_close=max(0.01, min(0.20, float(lp2_delta_close or 0.03))),
        prb1_delta_break=max(0.01, min(0.30, float(prb1_delta_break or 0.05))),
        prb1_delta_close=max(0.01, min(0.25, float(prb1_delta_close or 0.05))),
        prb2_delta_break=max(0.01, min(0.30, float(prb2_delta_break or 0.05))),
        prb2_delta_close=max(0.01, min(0.25, float(prb2_delta_close or 0.05))),
        lp1_volume_mult=max(0.5, min(3.0, float(lp1_volume_mult or 1.0))),
        lp2_volume_mult=max(0.5, min(3.0, float(lp2_volume_mult or 1.0))),
        prb_volume_mult=max(0.5, min(3.0, float(prb_volume_mult or 1.2))),
        range_cap=max(1.0, min(5.0, float(range_cap or 2.5))),
        two_bar_range_cap=max(1.5, min(8.0, float(two_bar_range_cap or 4.0))),
        stop_buffer_atr=max(0.05, min(0.50, float(stop_buffer_atr or 0.25))),
        stop_variant=selected_stop_variant,
        approach_lookback=max(2, min(10, int(approach_lookback or 4))),
        approach_min_atr=max(0.0, min(0.50, float(approach_min_atr or 0.05))),
        gap_max_atr=max(0.20, min(3.0, float(gap_max_atr or 1.2))),
        chase_max_atr=max(0.05, min(1.5, float(chase_max_atr or 0.30))),
        lp2_overextended_atr=max(0.50, min(3.0, float(lp2_overextended_atr or 1.2))),
        slippage_per_share=max(0.0, min(5.0, float(slippage_per_share or 0.0))),
        fees_per_share=max(0.0, min(5.0, float(fees_per_share or 0.0))),
        require_complete_quality=bool(require_complete_quality),
    )
    return run_market_screener(
        symbols=symbols,
        bars_loader=lambda symbol, selected_timeframe: _screener_frame(symbol, selected_timeframe, selected_asset),
        watchlist=watchlist if isinstance(watchlist, dict) else {},
        params=params,
    )


@app.get("/api/inefficiency-reclaim/settings")
def api_inefficiency_reclaim_settings():
    return {
        "min_daily_history_rows": SETTINGS.inefficiency_reclaim.min_daily_history_rows,
        "minimum": IRS_MIN_DAILY_HISTORY_ROWS_FLOOR,
        "maximum": IRS_MIN_DAILY_HISTORY_ROWS_CEILING,
        "env_key": IRS_MIN_DAILY_HISTORY_ROWS_KEY,
        "minimum_display_score": SETTINGS.inefficiency_reclaim.minimum_display_score,
        "display_score_minimum": IRS_MINIMUM_DISPLAY_SCORE_FLOOR,
        "display_score_maximum": SETTINGS.inefficiency_reclaim.minimum_order_score,
        "display_score_env_key": IRS_MINIMUM_DISPLAY_SCORE_KEY,
    }


@app.post("/api/inefficiency-reclaim/settings")
async def api_update_inefficiency_reclaim_settings(body: dict):
    try:
        if not isinstance(body, dict):
            raise ValueError("IRS settings payload must be an object.")
        has_rows = "min_daily_history_rows" in body
        has_score = "minimum_display_score" in body
        if not has_rows and not has_score:
            raise ValueError("No supported IRS setting was provided.")
        parsed_rows = (
            _parse_irs_min_daily_history_rows(body["min_daily_history_rows"])
            if has_rows
            else None
        )
        parsed_score = (
            _parse_irs_minimum_display_score(body["minimum_display_score"])
            if has_score
            else None
        )
        rows = (
            _set_irs_min_daily_history_rows(parsed_rows)
            if parsed_rows is not None
            else SETTINGS.inefficiency_reclaim.min_daily_history_rows
        )
        score = (
            _set_irs_minimum_display_score(parsed_score)
            if parsed_score is not None
            else SETTINGS.inefficiency_reclaim.minimum_display_score
        )
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    except OSError as exc:
        raise HTTPException(500, f"Could not save IRS settings: {exc}")
    saved = []
    if has_rows:
        saved.append(f"{IRS_MIN_DAILY_HISTORY_ROWS_KEY}={rows}")
    if has_score:
        saved.append(f"{IRS_MINIMUM_DISPLAY_SCORE_KEY}={score:g}")
    return {
        "ok": True,
        "min_daily_history_rows": rows,
        "minimum_display_score": score,
        "message": f"SAVED: {', '.join(saved)}",
    }


@app.get("/api/inefficiency-reclaim")
def api_inefficiency_reclaim(
    anchor_date: str = "",
    min_daily_history_rows: int | None = None,
):
    """Run the disabled-by-default IRS analysis scan over persisted bars."""
    from src.symbol_universe import load_stock_symbols

    try:
        config = SETTINGS.inefficiency_reclaim.strategy_config()
        parsed_anchor = str(anchor_date or "").strip()
        if parsed_anchor:
            datetime.strptime(parsed_anchor, "%Y-%m-%d")
        watchlist = _safe(da.load_watchlist, {})
        configured_symbols = _safe(load_stock_symbols, [])
        symbols = sorted(set(configured_symbols or []) | set((watchlist or {}).keys()))
        result = run_inefficiency_reclaim_screener(
            symbols=symbols,
            watchlist=watchlist if isinstance(watchlist, dict) else {},
            anchor_date=parsed_anchor or None,
            min_daily_history_rows=min_daily_history_rows,
            config=config,
            persist=False,
        )
        try:
            result["active_setups"] = InefficiencyReclaimStore(
                SETTINGS.inefficiency_reclaim.database_path
            ).active_setups()
        except Exception:
            result["active_setups"] = []
        return result
    except ValueError as exc:
        raise HTTPException(400, f"Invalid IRS scan request: {exc}")
    except Exception as exc:
        raise HTTPException(500, f"IRS scan failed closed: {exc}")


@app.get("/api/inefficiency-reclaim/active")
def api_inefficiency_reclaim_active():
    try:
        store = InefficiencyReclaimStore(SETTINGS.inefficiency_reclaim.database_path)
        return {
            "strategy_enabled": SETTINGS.inefficiency_reclaim.enabled,
            "paper_live_mode": "PAPER",
            "active_setups": store.active_setups(),
        }
    except Exception as exc:
        raise HTTPException(500, f"IRS state unavailable: {exc}")


@app.post("/api/watchlist/add")
async def api_watchlist_add(body: dict):
    from src.symbol_universe import add_stock_symbol, normalize_stock_symbol

    try:
        symbol = normalize_stock_symbol(str(body.get("symbol", "")))
    except ValueError as exc:
        raise HTTPException(400, str(exc))

    try:
        add_result = add_stock_symbol(symbol)
        if not add_result.get("added"):
            return {
                "ok": True,
                "symbol": symbol,
                "added": False,
                "duplicate": True,
                "pid": None,
                "log": None,
                "message": f"{symbol} is already in the stock universe. Onboarding was not queued again.",
            }
        SETTINGS.paths.runtime_dir.mkdir(parents=True, exist_ok=True)
        log_path = SETTINGS.paths.runtime_dir / f"onboard_symbol_{symbol}.log"
        client_id = _onboard_client_id(symbol)
        cmd = [
            sys.executable,
            "-m",
            "src.main",
            "--job",
            "onboard_symbol",
            "--symbol",
            symbol,
            "--client-id",
            client_id,
        ]
        with log_path.open("a", encoding="utf-8") as log:
            log.write(f"\n--- {datetime.now(ET).isoformat()} queued {' '.join(cmd)} ---\n")
            process = subprocess.Popen(
                cmd,
                cwd=str(ROOT),
                stdout=log,
                stderr=subprocess.STDOUT,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
        return {
            "ok": True,
            "symbol": symbol,
            "added": bool(add_result.get("added")),
            "pid": process.pid,
            "log": str(log_path),
            "message": f"Queued onboarding for {symbol}. The bot will fetch bars, calculate levels/zones, and refresh the working watchlist.",
        }
    except Exception as exc:
        raise HTTPException(500, str(exc))


@app.post("/api/watchlist/bulk-add")
async def api_watchlist_bulk_add(body: dict):
    from src.symbol_universe import add_stock_symbol, load_stock_symbols

    parsed, invalid = _parse_stock_symbol_upload(body)
    if not parsed and not invalid:
        raise HTTPException(400, "Upload a CSV containing comma-separated stock symbols.")
    existing = set(_safe(load_stock_symbols, []) or [])
    duplicates = [symbol for symbol in parsed if symbol in existing]
    new_symbols = [symbol for symbol in parsed if symbol not in existing]
    added: list[str] = []
    errors: list[dict[str, str]] = []
    for symbol in new_symbols:
        try:
            result = add_stock_symbol(symbol)
            if result.get("added"):
                added.append(symbol)
                existing.add(symbol)
            else:
                duplicates.append(symbol)
        except Exception as exc:
            errors.append({"symbol": symbol, "reason": str(exc)})
    queue = _queue_bulk_stock_onboarding(added) if added else {"pid": None, "log": None, "queue_file": None}
    message_parts = []
    if added:
        message_parts.append(f"Queued bulk onboarding for {len(added)} new symbol{'s' if len(added) != 1 else ''}, followed by a premarket refresh.")
    if duplicates:
        message_parts.append(f"Skipped {len(duplicates)} duplicate{'s' if len(duplicates) != 1 else ''}.")
    if invalid:
        message_parts.append(f"Ignored {len(invalid)} invalid entr{'ies' if len(invalid) != 1 else 'y'}.")
    if errors:
        message_parts.append(f"{len(errors)} symbol{'s' if len(errors) != 1 else ''} failed to add.")
    return {
        "ok": not errors,
        "requested": len(parsed) + len(invalid),
        "parsed": parsed,
        "added": added,
        "duplicates": sorted(set(duplicates)),
        "invalid": invalid,
        "errors": errors,
        "pid": queue.get("pid"),
        "log": queue.get("log"),
        "queue_file": queue.get("queue_file"),
        "message": " ".join(message_parts) or "No new symbols to onboard.",
    }


@app.post("/api/watchlist/remove")
async def api_watchlist_remove(body: dict):
    from src.symbol_universe import normalize_stock_symbol, remove_stock_symbol

    try:
        symbol = normalize_stock_symbol(str(body.get("symbol", "")))
    except ValueError as exc:
        raise HTTPException(400, str(exc))

    try:
        remove_result = remove_stock_symbol(symbol)
        state_path = SETTINGS.paths.state_file
        removed_from_state = False
        if state_path.exists():
            try:
                state = json.loads(state_path.read_text(encoding="utf-8"))
            except (json.JSONDecodeError, OSError):
                state = {}
            if isinstance(state, dict):
                watchlist = state.get("watchlist")
                if isinstance(watchlist, dict) and symbol in watchlist:
                    watchlist.pop(symbol, None)
                    removed_from_state = True
                temp_path = state_path.with_suffix(state_path.suffix + ".tmp")
                temp_path.write_text(json.dumps(state, indent=2), encoding="utf-8")
                os.replace(temp_path, state_path)
        return {
            "ok": True,
            "symbol": symbol,
            "removed": bool(remove_result.get("removed")) or removed_from_state,
            "removed_from_config": bool(remove_result.get("removed")),
            "removed_from_state": removed_from_state,
            "message": f"{symbol} removed from stock universe and active watchlist.",
        }
    except Exception as exc:
        raise HTTPException(500, str(exc))


def _strategy_payload(mode: str = "bmsb", asset: str = "stock") -> dict:
    from src.symbol_universe import load_stock_symbols, stock_symbols_file

    selected_asset = "crypto" if str(asset).strip().lower() == "crypto" else "stock"
    if selected_asset == "crypto":
        index = _safe(da.crypto_bars_index, {})
        configured = _safe(da.load_crypto_symbols, [])
        configured_symbols = [
            str(row.get("inst_id") or "").upper()
            for row in configured
            if isinstance(row, dict) and row.get("inst_id")
        ]
        symbols = (
            sorted(set(configured_symbols))
            if CRYPTO_SETTINGS.symbols_file.exists()
            else sorted(set(index or {}))
        )
    else:
        index = _safe(da.bars_index, {})
        configured_symbols = _safe(load_stock_symbols, [])
        if stock_symbols_file().exists():
            symbols = sorted(configured_symbols or [])
        else:
            symbols = sorted(
                symbol for symbol, frames in (index or {}).items()
                if isinstance(frames, dict) and ("daily" in frames or "weekly" in frames)
            )
    today = datetime.now(ET).date()
    selected_mode = str(mode or "bmsb").strip().lower()
    if selected_mode not in {"bmsb", "gaussian"}:
        selected_mode = "bmsb"
    scanner = _gaussian_scan_symbol if selected_mode == "gaussian" else _bmsb_scan_symbol
    strategy_name = "Gaussian Channel Strategy" if selected_mode == "gaussian" else "BMSB Strategy"
    threshold_pct = GAUSSIAN_NEAR_THRESHOLD_PCT if selected_mode == "gaussian" else BMSB_NEAR_CROSS_THRESHOLD_PCT
    recent_days = GAUSSIAN_RECENT_SIGNAL_DAYS if selected_mode == "gaussian" else BMSB_RECENT_CROSS_DAYS
    description = (
        "Gaussian Channel monitor. Signals are based on bullish close crosses over the upper channel or symbols near that trigger."
        if selected_mode == "gaussian"
        else "Weekly BMSB monitor. Signals are based on the 21W EMA crossing the 20W SMA, or symbols near that cross."
    )
    rows = []
    for symbol in symbols:
        row = _safe(
            lambda s=symbol: scanner(s, today, include_neutral=True, asset=selected_asset),
            None,
        )
        if row is None:
            row = {
                "symbol": symbol,
                "signal": "NO_DATA",
                "side": "NONE",
                "urgency": "none",
                "last_price": None,
                "ema21": None,
                "sma20": None,
                "gap": None,
                "gap_pct": None,
                "entry_level": None,
                "exit_level": None,
                "trigger_level": None,
                "trigger_gap_pct": None,
                "cross_date": None,
                "daily_bar": None,
                "weekly_bar": None,
                "threshold_pct": threshold_pct,
            }
        rows.append(row)
    matched_rows = [row for row in rows if row.get("urgency") in {"crossed", "near"}]
    urgency_rank = {"crossed": 0, "near": 1}
    side_rank = {"LONG": 0, "EXIT": 1, "NONE": 2}
    def _sort_gap(row: dict) -> float:
        value = row.get("trigger_gap_pct")
        if value is None:
            value = row.get("gap_pct")
        try:
            return float(value)
        except (TypeError, ValueError):
            return 999.0
    rows.sort(key=lambda r: (
        urgency_rank.get(str(r.get("urgency")), 9),
        side_rank.get(str(r.get("side")), 9),
        _sort_gap(r),
        str(r.get("symbol", "")),
    ))
    return {
        "asset": selected_asset,
        "mode": selected_mode,
        "strategy": strategy_name,
        "description": description,
        "threshold_pct": threshold_pct,
        "recent_days": recent_days,
        "symbols_scanned": len(symbols),
        "symbols": symbols,
        "matches": len(matched_rows),
        "crossed": sum(1 for row in matched_rows if row.get("urgency") == "crossed"),
        "near": sum(1 for row in matched_rows if row.get("urgency") == "near"),
        "rows": rows,
    }


@app.get("/api/strategy")
def api_strategy(mode: str = "bmsb"):
    return _strategy_payload(mode, "stock")


@app.get("/api/crypto/strategy")
def api_crypto_strategy(mode: str = "bmsb"):
    return _strategy_payload(mode, "crypto")


@app.get("/api/preopen")
def api_preopen():
    watchlist = _safe(da.load_watchlist, {})
    rows = []
    for sym, plan in sorted((watchlist or {}).items()):
        if not isinstance(plan, dict): continue
        spacing = plan.get("level_spacing", {}) or {}
        levels = plan.get("levels", [])
        blocked = bool(plan.get("news_blocked"))
        rows.append({
            "symbol": sym,
            "status": "BLOCKED" if blocked else ("READY" if levels else "NO LEVELS"),
            "price": spacing.get("current_price"),
            "daily_atr": plan.get("daily_atr"),
            "technical_atr": plan.get("technical_atr"),
            "trade_levels": len(levels),
            "room": spacing.get("reason", "-"),
            "news": "BLOCKED" if blocked else "CLEAR",
        })
    return {"rows": rows, "watchlist": watchlist}


@app.get("/api/preopen/{symbol}")
def api_preopen_symbol(symbol: str):
    symbol = symbol.upper()
    watchlist = _safe(da.load_watchlist, {})
    plan = (watchlist or {}).get(symbol)
    if plan is None:
        raise HTTPException(404, f"{symbol} not in watchlist")
    spacing = plan.get("level_spacing", {}) or {}
    raw_levels = plan.get("raw_levels", [])
    trade_levels = plan.get("levels", [])
    return {
        "symbol": symbol,
        "price": spacing.get("current_price"),
        "daily_atr": plan.get("daily_atr"),
        "technical_atr": plan.get("technical_atr"),
        "trade_levels": trade_levels,
        "raw_levels": raw_levels,
        "room": spacing.get("reason", "-"),
        "news": "BLOCKED" if plan.get("news_blocked") else "CLEAR",
        "headlines": plan.get("matched_headlines", []),
        "security_type": plan.get("security_type", "-"),
    }


@app.get("/api/bars/{symbol}/{timeframe}")
def api_bars(symbol: str, timeframe: str):
    symbol = symbol.upper()

    def _fifteen_minute_bars() -> pd.DataFrame:
        stored = da.get_bars(symbol, "intraday_15m")
        if not stored.empty:
            return stored
        intraday = da.get_bars(symbol, "intraday_5m")
        if intraday.empty:
            return intraday
        frame = intraday.set_index("date").sort_index()
        return (
            frame.resample("15min", origin="start_day", offset="30min")
            .agg({"open": "first", "high": "max", "low": "min", "close": "last", "volume": "sum"})
            .dropna(subset=["open", "high", "low", "close"])
            .reset_index()
        )

    def _thirty_minute_bars() -> pd.DataFrame:
        stored = da.get_bars(symbol, "intraday_30m")
        if not stored.empty:
            return stored
        intraday = da.get_bars(symbol, "intraday_5m")
        if intraday.empty:
            return intraday
        frame = intraday.set_index("date").sort_index()
        return (
            frame.resample("30min", origin="start_day", offset="30min")
            .agg({"open": "first", "high": "max", "low": "min", "close": "last", "volume": "sum"})
            .dropna(subset=["open", "high", "low", "close"])
            .reset_index()
        )

    def _hourly_bars() -> pd.DataFrame:
        intraday = da.get_bars(symbol, "intraday_5m")
        if intraday.empty:
            return intraday
        frame = intraday.set_index("date").sort_index()
        return (
            frame.resample("60min", origin="start_day", offset="30min")
            .agg({"open": "first", "high": "max", "low": "min", "close": "last", "volume": "sum"})
            .dropna(subset=["open", "high", "low", "close"])
            .reset_index()
        )

    def _four_hour_bars() -> pd.DataFrame:
        stored = da.get_bars(symbol, "intraday_4h")
        if not stored.empty:
            return stored
        intraday = da.get_bars(symbol, "intraday_5m")
        if intraday.empty:
            return intraday
        frame = intraday.set_index("date").sort_index()
        return (
            frame.resample("240min", origin="start_day", offset="30min")
            .agg({"open": "first", "high": "max", "low": "min", "close": "last", "volume": "sum"})
            .dropna(subset=["open", "high", "low", "close"])
            .reset_index()
        )

    def _weekly_bars() -> pd.DataFrame:
        stored = da.get_bars(symbol, "weekly")
        if not stored.empty:
            return stored
        daily = da.get_bars(symbol, "daily")
        if daily.empty:
            return daily
        frame = daily.set_index("date").sort_index()
        return (
            frame.resample("W-FRI")
            .agg({"open": "first", "high": "max", "low": "min", "close": "last", "volume": "sum"})
            .dropna(subset=["open", "high", "low", "close"])
            .reset_index()
        )

    def _live_daily_bars() -> pd.DataFrame:
        daily = da.get_bars(symbol, "daily")
        intraday = da.get_bars(symbol, "intraday_5m")
        if intraday.empty:
            return daily

        derived_daily = daily_bars_from_intraday(intraday)
        if derived_daily.empty:
            return daily
        if daily.empty:
            return derived_daily

        derived_dates = pd.to_datetime(derived_daily["date"], errors="coerce").dt.date
        historical_dates = set(pd.to_datetime(daily["date"], errors="coerce").dropna().dt.date)
        missing_or_live = derived_daily.loc[~derived_dates.isin(historical_dates)]
        if missing_or_live.empty:
            return daily
        return pd.concat([daily, missing_or_live], ignore_index=True).sort_values("date")

    if timeframe == "intraday_15m":
        df = _safe(_fifteen_minute_bars, None)
    elif timeframe == "intraday_30m":
        df = _safe(_thirty_minute_bars, None)
    elif timeframe == "intraday_1h":
        df = _safe(_hourly_bars, None)
    elif timeframe == "intraday_4h":
        df = _safe(_four_hour_bars, None)
    elif timeframe == "daily_live":
        df = _safe(_live_daily_bars, None)
    elif timeframe == "weekly":
        df = _safe(_weekly_bars, None)
    else:
        df = _safe(lambda: da.get_bars(symbol, timeframe), None)
    if df is None or df.empty:
        return {"bars": [], "stale": True, "last_date": None}
    chart_limits = {
        "intraday_5m": 500,
        "intraday_15m": 700,
        "intraday_30m": 700,
        "intraday_1h": 500,
        "intraday_4h": 500,
        "daily_live": 600,
        "daily": 600,
        "weekly": 260,
    }
    df = df.tail(chart_limits.get(timeframe, 300))
    last_date = str(df["date"].iloc[-1]) if "date" in df.columns else None
    # stale = last bar is genuinely missing sessions.
    # Count how many weekdays (Mon-Fri) lie between last_bar and today — if more
    # than 1 then at least one trading session was skipped (we can't know about
    # holidays without a full calendar, so we allow 1 extra day as holiday buffer).
    stale = False
    if last_date:
        from datetime import date, timedelta
        try:
            last_dt = date.fromisoformat(str(last_date)[:10])
            today = datetime.now(ET).date()
            # count weekdays strictly between last_dt and today (exclusive of both)
            weekdays_gap = 0
            d = last_dt + timedelta(days=1)
            while d < today:
                if d.weekday() < 5:
                    weekdays_gap += 1
                d += timedelta(days=1)
            # Weekly bars naturally lag within the current week.
            stale = weekdays_gap > (7 if timeframe == "weekly" else 1)
        except Exception:
            pass
    return {
        "bars": df.rename(columns={"date": "t"}).to_dict(orient="records"),
        "stale": stale,
        "last_date": last_date,
    }


def _crypto_tv_payload(tv_state: dict, *, limit: int = 50) -> dict:
    if not isinstance(tv_state, dict):
        tv_state = {}
    positions = tv_state.get("positions", {})
    executions = tv_state.get("executions", [])
    if not isinstance(positions, dict):
        positions = {}
    if not isinstance(executions, list):
        executions = []
    return {
        "updated_at": tv_state.get("updated_at"),
        "positions": list(positions.values()),
        "executions": [
            {key: value for key, value in row.items() if key not in {"okx_response", "raw"}}
            for row in executions[:limit]
            if isinstance(row, dict)
        ],
    }


def _crypto_event_date(value: Any):
    if not value:
        return None
    try:
        text = str(value).replace("Z", "+00:00")
        parsed = datetime.fromisoformat(text)
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=ZoneInfo("UTC"))
        return parsed.astimezone(ET).date()
    except Exception:
        return None


def _resolve_crypto_add_instrument(raw_symbol: str) -> tuple[str, list[str]]:
    """Resolve user-friendly crypto input to an OKX instrument id.

    People often type TradingView-style pairs such as DOGE-USD while OKX spot
    commonly lists DOGE-USDT. Prefer an exact OKX match, then try the USDT spot
    proxy for plain USD pairs.
    """

    normalized = normalize_okx_instrument(raw_symbol)
    if not normalized:
        raise ValueError("Crypto symbol is required.")

    candidates = [normalized]
    if normalized.endswith("-USD") and not normalized.endswith("-USDT"):
        candidates.append(f"{normalized[:-4]}-USDT")

    by_type: dict[str, set[str]] = {}
    client = OKXClient()
    for candidate in candidates:
        by_type.setdefault(default_instrument_type(candidate), set()).add(candidate)

    available: set[str] = set()
    for inst_type in by_type:
        available.update(item.inst_id for item in client.instruments(inst_type))

    for candidate in candidates:
        if candidate in available:
            return candidate, candidates

    raise ValueError(
        f"{normalized} is not available on OKX. Tried: {', '.join(candidates)}."
    )


@app.get("/api/crypto")
def api_crypto():
    state = _safe(da.load_crypto_dashboard_state, {})
    tv_state = _safe(load_tradingview_state, {})
    configured = _safe(da.load_crypto_symbols, [])
    watchlist = state.get("watchlist", {}) if isinstance(state, dict) else {}
    index = _safe(da.crypto_bars_index, {})
    tv_positions = tv_state.get("positions", {}) if isinstance(tv_state, dict) else {}
    tv_executions = tv_state.get("executions", []) if isinstance(tv_state, dict) else []
    if not isinstance(tv_positions, dict):
        tv_positions = {}
    if not isinstance(tv_executions, list):
        tv_executions = []
    today = datetime.now(ET).date()
    position_by_symbol = {}
    for inst_id, position in tv_positions.items():
        if not isinstance(position, dict):
            continue
        try:
            qty = float(position.get("qty") or 0)
        except (TypeError, ValueError):
            qty = 0.0
        if qty > 0:
            position_by_symbol[str(inst_id).upper()] = {**position, "qty": qty}
    realized_pnl_by_symbol = {}
    for execution in tv_executions:
        if not isinstance(execution, dict):
            continue
        if str(execution.get("status", "")).lower() != "submitted":
            continue
        if str(execution.get("action", "")).upper() != "SELL":
            continue
        if _crypto_event_date(execution.get("created_at")) != today:
            continue
        inst_id = str(execution.get("inst_id") or "").upper()
        if not inst_id:
            continue
        try:
            pnl = float(execution.get("daily_pnl") if execution.get("daily_pnl") is not None else execution.get("estimated_pnl") or 0)
        except (TypeError, ValueError):
            pnl = 0.0
        realized_pnl_by_symbol[inst_id] = realized_pnl_by_symbol.get(inst_id, 0.0) + pnl
    rows = []
    configured_by_inst = {row.get("inst_id"): row for row in configured if isinstance(row, dict)}
    symbols = sorted(set(configured_by_inst) | set(watchlist or {}) | set(index or {}))
    for symbol in symbols:
        plan = watchlist.get(symbol, {}) if isinstance(watchlist, dict) else {}
        frames = index.get(symbol, {}) if isinstance(index, dict) else {}
        trade_levels = plan.get("trade_levels", []) if isinstance(plan, dict) else []
        raw_levels = plan.get("raw_levels", []) if isinstance(plan, dict) else []
        latest = None
        for tf in ("intraday_5m", "intraday_1h", "intraday_4h", "daily"):
            latest = (frames.get(tf, {}) or {}).get("last_bar") if isinstance(frames.get(tf, {}), dict) else latest
            if latest:
                break
        price = plan.get("price") if isinstance(plan, dict) else None
        daily_open = None
        daily_change = None
        daily_change_pct = None
        daily_pnl = None
        daily_pnl_pct = None
        daily_pnl_status = None
        position = position_by_symbol.get(symbol.upper())
        position_qty = float((position or {}).get("qty") or 0.0)
        daily_bars = _safe(lambda s=symbol: da.get_crypto_bars(s, "daily"), None)
        try:
            current_price = float(price) if price is not None else None
        except (TypeError, ValueError):
            current_price = None
        if daily_bars is not None and not daily_bars.empty:
            try:
                last_daily = daily_bars.iloc[-1]
                daily_open = float(last_daily.get("open"))
                latest_daily_close = float(last_daily.get("close"))
                current_price = float(price) if price is not None else latest_daily_close
                if daily_open > 0:
                    daily_change = current_price - daily_open
                    daily_change_pct = (daily_change / daily_open) * 100.0
            except (TypeError, ValueError):
                pass
        realized_today = realized_pnl_by_symbol.get(symbol.upper())
        has_pnl_today = realized_today is not None
        if has_pnl_today:
            daily_pnl = float(realized_today or 0.0)
            daily_pnl_status = "CLOSED"
        if position_qty > 0 and current_price is not None:
            try:
                opened_today = _crypto_event_date(position.get("opened_at")) == today if isinstance(position, dict) else False
                entry = float(position.get("entry") or 0.0) if isinstance(position, dict) else 0.0
                baseline = entry if opened_today and entry > 0 else daily_open
                if baseline and baseline > 0:
                    open_pnl = (current_price - baseline) * position_qty
                    daily_pnl = (daily_pnl or 0.0) + open_pnl
                    daily_pnl_pct = ((current_price - baseline) / baseline) * 100.0
                    daily_pnl_status = "OPEN"
            except (TypeError, ValueError):
                pass
        rows.append({
            "symbol": symbol,
            "raw": configured_by_inst.get(symbol, {}).get("raw"),
            "inst_type": configured_by_inst.get(symbol, {}).get("inst_type"),
            "ready": bool(plan.get("ready")) if isinstance(plan, dict) else False,
            "price": price,
            "daily_atr": plan.get("daily_atr") if isinstance(plan, dict) else None,
            "daily_open": daily_open,
            "daily_change": daily_change,
            "daily_change_pct": daily_change_pct,
            "daily_pnl": daily_pnl,
            "daily_pnl_pct": daily_pnl_pct,
            "daily_pnl_status": daily_pnl_status,
            "position_qty": position_qty,
            "trade_levels": len(trade_levels),
            "raw_levels": len(raw_levels),
            "last_bar": latest,
            "bars": plan.get("bars", {}) if isinstance(plan, dict) else {},
            "reason": plan.get("reason") if isinstance(plan, dict) else None,
        })
    ready = sum(1 for row in rows if row.get("ready"))
    return {
        "updated_at": state.get("updated_at") if isinstance(state, dict) else None,
        "symbols": symbols,
        "rows": rows,
        "errors": state.get("errors", []) if isinstance(state, dict) else [],
        "tradingview": _crypto_tv_payload(tv_state, limit=50),
        "metrics": {
            "configured": len(configured),
            "tracked": len(rows),
            "ready": ready,
            "with_trade_levels": sum(1 for row in rows if int(row.get("trade_levels") or 0) > 0),
        },
    }


@app.get("/api/crypto/tradingview")
def api_crypto_tradingview_state():
    tv_state = _safe(load_tradingview_state, {})
    return _crypto_tv_payload(tv_state, limit=100)


@app.post("/api/crypto/add")
async def api_crypto_add(body: dict):
    try:
        raw_symbol = str((body or {}).get("symbol", "") or "").strip()
        normalized = normalize_okx_instrument(raw_symbol)
        if not normalized:
            raise ValueError("Crypto symbol is required.")
        duplicate_candidates = [normalized]
        if normalized.endswith("-USD") and not normalized.endswith("-USDT"):
            duplicate_candidates.append(f"{normalized[:-4]}-USDT")
        configured_ids = {
            str(row.get("inst_id") or "").upper()
            for row in _safe(da.load_crypto_symbols, [])
            if isinstance(row, dict) and row.get("inst_id")
        }
        duplicate = next((candidate for candidate in duplicate_candidates if candidate in configured_ids), None)
        if duplicate:
            return {
                "ok": True,
                "symbol": duplicate,
                "added": False,
                "duplicate": True,
                "pid": None,
                "log": None,
                "message": f"{duplicate} is already in the crypto universe. Candle loading was not queued again.",
            }

        inst_id, candidates = _resolve_crypto_add_instrument(raw_symbol)
        add_result = add_crypto_symbol(inst_id)
        if not add_result.get("added"):
            return {
                "ok": True,
                "symbol": inst_id,
                "added": False,
                "duplicate": True,
                "pid": None,
                "log": None,
                "message": f"{inst_id} is already in the crypto universe. Candle loading was not queued again.",
            }

        runtime_dir = CRYPTO_SETTINGS.memory_dir / "runtime"
        runtime_dir.mkdir(parents=True, exist_ok=True)
        safe_name = "".join(char if char.isalnum() else "_" for char in inst_id)
        log_path = runtime_dir / f"onboard_crypto_{safe_name}.log"
        cmd = [
            sys.executable,
            "-m",
            "src.crypto.main",
            "--job",
            "collect_analyze",
            "--symbol",
            inst_id,
        ]
        with log_path.open("a", encoding="utf-8") as log:
            log.write(f"\n--- {datetime.now(ET).isoformat()} queued {' '.join(cmd)} ---\n")
            if len(candidates) > 1 and inst_id != candidates[0]:
                log.write(f"Resolved requested symbol {raw_symbol} to OKX instrument {inst_id}.\n")
            process = subprocess.Popen(
                cmd,
                cwd=str(ROOT),
                stdout=log,
                stderr=subprocess.STDOUT,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
        return {
            "ok": True,
            "symbol": inst_id,
            "requested": raw_symbol,
            "added": True,
            "duplicate": False,
            "pid": process.pid,
            "log": str(log_path),
            "message": f"Queued crypto onboarding for {inst_id}. The bot will fetch candles, calculate levels/zones, and refresh the Crypto tab.",
        }
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    except Exception as exc:
        raise HTTPException(500, str(exc))


@app.post("/api/crypto/bulk-add")
async def api_crypto_bulk_add(body: dict):
    raw_text = str((body or {}).get("csv_text", "") or "")
    values = [
        value.strip().strip('"\'')
        for value in re.split(r"[\s,;]+", raw_text)
        if value.strip().strip('"\'')
    ]
    values = [value for value in values if value.upper() not in {"SYMBOL", "INSTRUMENT", "INST_ID"}]
    if not values:
        raise HTTPException(400, "Choose a CSV file or provide crypto symbols first.")

    added: list[str] = []
    duplicates: list[str] = []
    invalid: list[dict[str, str]] = []
    pids: list[int] = []
    seen: set[str] = set()
    for value in values:
        normalized = normalize_okx_instrument(value)
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        try:
            result = await api_crypto_add({"symbol": value})
        except HTTPException as exc:
            invalid.append({"value": value, "reason": str(exc.detail)})
            continue
        symbol = str(result.get("symbol") or normalized)
        if result.get("duplicate"):
            duplicates.append(symbol)
        elif result.get("added"):
            added.append(symbol)
            if result.get("pid"):
                pids.append(int(result["pid"]))

    return {
        "ok": True,
        "added": added,
        "duplicates": duplicates,
        "invalid": invalid,
        "pids": pids,
        "message": f"Queued {len(added)} crypto symbol(s); {len(duplicates)} duplicate(s); {len(invalid)} invalid.",
    }


@app.post("/api/crypto/remove")
async def api_crypto_remove(body: dict):
    try:
        symbol = normalize_okx_instrument(str((body or {}).get("symbol", "") or ""))
        if not symbol:
            raise ValueError("Crypto symbol is required.")
        result = remove_crypto_symbol(symbol)
        state_path = CRYPTO_SETTINGS.memory_dir / "state.json"
        removed_from_state = False
        if state_path.exists():
            state = _safe(lambda: json.loads(state_path.read_text(encoding="utf-8")), {})
            if isinstance(state, dict):
                symbols = state.get("symbols")
                if isinstance(symbols, list):
                    filtered = [item for item in symbols if normalize_okx_instrument(str(item)) != symbol]
                    removed_from_state = len(filtered) != len(symbols)
                    state["symbols"] = filtered
                watchlist = state.get("watchlist")
                if isinstance(watchlist, dict) and symbol in watchlist:
                    watchlist.pop(symbol, None)
                    removed_from_state = True
                temp_path = state_path.with_suffix(state_path.suffix + ".tmp")
                temp_path.write_text(json.dumps(state, indent=2), encoding="utf-8")
                os.replace(temp_path, state_path)
        return {
            "ok": True,
            "symbol": symbol,
            "removed": bool(result.get("removed")) or removed_from_state,
            "removed_from_config": bool(result.get("removed")),
            "removed_from_state": removed_from_state,
            "message": f"{symbol} removed from the crypto universe and active monitor. Saved candle files were kept.",
        }
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    except Exception as exc:
        raise HTTPException(500, str(exc))


@app.post("/api/crypto/order/simulate")
async def api_crypto_order_simulate(body: dict):
    try:
        symbol = normalize_okx_instrument(str((body or {}).get("symbol", "") or ""))
        if not symbol:
            raise ValueError("Crypto symbol is required.")
        side = "SELL" if str((body or {}).get("signal", "")).upper() == "SELL" else "BUY"
        cash = float((body or {}).get("available_cash") or 0.0)
        request = CryptoOrderRequest(
            symbol=symbol,
            side=side,
            entry=float((body or {}).get("entry") or 0.0),
            stop=float((body or {}).get("stop") or 0.0),
            target=float((body or {}).get("target")) if (body or {}).get("target") is not None else None,
            quantity=float((body or {}).get("quantity") or 0.0) or None,
            account_equity=cash,
            cash_available=cash,
        )
        result = CryptoOrderManager().place_order(request, live=False)
        plan = result.get("plan") or result
        return {
            "ok": bool(result.get("ok")),
            "message": (
                f"Simulated OKX demo order planned for {symbol}. No order was submitted."
                if result.get("ok")
                else f"Crypto order rejected: {', '.join(plan.get('reasons') or ['invalid setup'])}."
            ),
            "id": None,
            "dry_run": True,
            "paper": True,
            "plan": plan,
        }
    except (TypeError, ValueError) as exc:
        raise HTTPException(400, str(exc))
    except Exception as exc:
        raise HTTPException(500, str(exc))


@app.get("/api/crypto/orders/manual")
def api_crypto_manual_orders():
    state = load_manual_order_state()
    return {
        "ok": True,
        "updated_at": state.get("updated_at"),
        "orders": state.get("orders", []),
        "paper": True,
        "demo": True,
        "live_enabled": False,
        "credentials_configured": bool(
            CRYPTO_SETTINGS.okx_api_key
            and CRYPTO_SETTINGS.okx_api_secret
            and CRYPTO_SETTINGS.okx_api_passphrase
        ),
        "simulated_trading": bool(CRYPTO_SETTINGS.okx_simulated_trading),
        "account_mode": CRYPTO_SETTINGS.okx_account_mode,
    }


@app.post("/api/crypto/order/place")
async def api_crypto_order_place(body: dict):
    """Submit a manual Strategy Crypto order to OKX Demo only."""
    try:
        body = body or {}
        symbol = normalize_okx_instrument(str((body or {}).get("symbol", "") or ""))
        if not symbol:
            raise ValueError("Crypto symbol is required.")
        side = "SELL" if str(body.get("signal", "")).upper() == "SELL" else "BUY"
        cash = float(body.get("available_cash") or 0.0)
        max_risk_value = body.get("max_open_risk_pct")
        max_risk_pct = None
        if max_risk_value not in (None, ""):
            max_risk_percent = float(max_risk_value)
            if not math.isfinite(max_risk_percent) or max_risk_percent <= 0 or max_risk_percent > 100:
                raise ValueError("Max risk % must be greater than 0 and no more than 100.")
            max_risk_pct = max_risk_percent / 100.0
        reward_risk_value = body.get("reward_risk")
        min_reward_risk_ratio = None
        if reward_risk_value not in (None, ""):
            min_reward_risk_ratio = float(reward_risk_value)
            if (
                not math.isfinite(min_reward_risk_ratio)
                or min_reward_risk_ratio <= 0
                or min_reward_risk_ratio > 100
            ):
                raise ValueError("Reward:risk must be greater than 0 and no more than 100.")
        request = CryptoOrderRequest(
            symbol=symbol,
            side=side,
            entry=float(body.get("entry") or 0.0),
            stop=float(body.get("stop") or 0.0),
            target=float(body.get("target")) if body.get("target") is not None else None,
            quantity=float(body.get("quantity") or 0.0) or None,
            account_equity=cash,
            cash_available=cash,
            max_risk_pct=max_risk_pct,
            min_reward_risk_ratio=min_reward_risk_ratio,
        )
        result = execute_manual_demo_order(
            request,
            order_type=str(body.get("entry_order_type") or "LIMIT"),
        )
        return {
            **result,
            "paper": True,
            "demo": True,
            "dry_run": False,
            "live": False,
        }
    except (TypeError, ValueError) as exc:
        raise HTTPException(400, str(exc))
    except Exception as exc:
        raise HTTPException(500, str(exc))


@app.post("/api/crypto/tradingview")
@app.post("/api/webhook/tradingview/crypto")
async def api_crypto_tradingview_webhook(body: dict, background_tasks: BackgroundTasks):
    try:
        accepted = enqueue_tradingview_webhook(body or {})
        execution_id = str((accepted.get("execution") or {}).get("id") or "")
        if execution_id and accepted.get("queued"):
            background_tasks.add_task(process_queued_tradingview_execution, execution_id)
        return accepted
    except PermissionError as exc:
        raise HTTPException(403, str(exc))
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    except Exception as exc:
        raise HTTPException(500, str(exc))


@app.get("/api/stocks/tradingview")
@app.get("/api/webhook/tradingview/stocks")
def api_stock_tradingview_state():
    state = _safe(load_stock_tradingview_state, {})
    executions = state.get("executions", []) if isinstance(state, dict) else []
    return {
        "updated_at": state.get("updated_at") if isinstance(state, dict) else None,
        "executions": executions[:100] if isinstance(executions, list) else [],
    }


@app.post("/api/stocks/tradingview")
@app.post("/api/webhook/tradingview/stocks")
async def api_stock_tradingview_webhook(body: dict):
    try:
        return handle_stock_tradingview_webhook(body or {})
    except PermissionError as exc:
        raise HTTPException(403, str(exc))
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    except Exception as exc:
        raise HTTPException(500, str(exc))


def _forex_request_rows(limit: int = 30) -> list[dict]:
    requests = _safe(da.load_order_requests, [])
    rows = []
    for request in requests:
        if not isinstance(request, dict):
            continue
        symbol = str(request.get("symbol") or "")
        if str(request.get("source") or "") == "forex_tradingview" or fx_pair_components(symbol):
            rows.append(request)
    return rows[:limit]


@app.get("/api/forex")
def api_forex():
    state = _safe(load_forex_tradingview_state, {})
    executions = state.get("executions", []) if isinstance(state, dict) else []
    return {
        "updated_at": state.get("updated_at") if isinstance(state, dict) else None,
        "symbols": sorted({str(item).upper() for item in SETTINGS.fx_symbols if str(item).strip()}),
        "webhook_url": "/api/webhook/tradingview/forex",
        "executions": executions[:100] if isinstance(executions, list) else [],
        "order_requests": _forex_request_rows(50),
    }


@app.get("/api/forex/tradingview")
@app.get("/api/webhook/tradingview/forex")
def api_forex_tradingview_state():
    state = _safe(load_forex_tradingview_state, {})
    executions = state.get("executions", []) if isinstance(state, dict) else []
    return {
        "updated_at": state.get("updated_at") if isinstance(state, dict) else None,
        "executions": executions[:100] if isinstance(executions, list) else [],
    }


@app.post("/api/forex/tradingview")
@app.post("/api/webhook/tradingview/forex")
async def api_forex_tradingview_webhook(body: dict):
    try:
        return handle_forex_tradingview_webhook(body or {})
    except PermissionError as exc:
        raise HTTPException(403, str(exc))
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    except Exception as exc:
        raise HTTPException(500, str(exc))


@app.get("/api/crypto/{symbol}")
def api_crypto_symbol(symbol: str):
    inst_id = normalize_okx_instrument(symbol)
    state = _safe(da.load_crypto_dashboard_state, {})
    watchlist = state.get("watchlist", {}) if isinstance(state, dict) else {}
    plan = watchlist.get(inst_id)
    if not isinstance(plan, dict):
        raise HTTPException(404, f"{inst_id} is not analyzed yet")
    return plan


@app.get("/api/crypto/bars/{symbol}/{timeframe}")
def api_crypto_bars(symbol: str, timeframe: str):
    inst_id = normalize_okx_instrument(symbol)
    tf = {
        "daily_live": "daily",
        "intraday_4h": "intraday_4h",
        "intraday_1h": "intraday_1h",
        "intraday_30m": "intraday_30m",
        "intraday_15m": "intraday_15m",
        "intraday_5m": "intraday_5m",
        "daily": "daily",
    }.get(timeframe, timeframe)
    df = _safe(lambda: da.get_crypto_bars(inst_id, tf), None)
    if (df is None or df.empty) and tf in {"intraday_15m", "intraday_30m", "intraday_1h", "intraday_4h"}:
        source = _safe(lambda: da.get_crypto_bars(inst_id, "intraday_5m"), None)
        if source is not None and not source.empty:
            minutes = {
                "intraday_15m": 15,
                "intraday_30m": 30,
                "intraday_1h": 60,
                "intraday_4h": 240,
            }[tf]
            frame = source.set_index("date").sort_index()
            df = (
                frame.resample(f"{minutes}min", origin="start_day")
                .agg({"open": "first", "high": "max", "low": "min", "close": "last", "volume": "sum"})
                .dropna(subset=["open", "high", "low", "close"])
                .reset_index()
            )
    if df is None or df.empty:
        return {"bars": [], "stale": True, "last_date": None}
    limits = {
        "intraday_5m": 500,
        "intraday_15m": 700,
        "intraday_30m": 700,
        "intraday_1h": 500,
        "intraday_4h": 500,
        "daily": 600,
        "daily_live": 600,
    }
    df = df.tail(limits.get(tf, 300))
    last_date = str(df["date"].iloc[-1]) if "date" in df.columns else None
    return {
        "bars": df.rename(columns={"date": "t"}).to_dict(orient="records"),
        "stale": False,
        "last_date": last_date,
    }


@app.post("/api/crypto/analyze")
async def api_crypto_analyze(body: dict | None = None):
    try:
        symbols = None
        if body and body.get("symbols"):
            raw = body.get("symbols")
            symbols = raw if isinstance(raw, list) else [raw]
        result = run_crypto_analysis(symbols)
        return {
            "ok": True,
            "updated_at": result.get("updated_at"),
            "symbols": result.get("symbols", []),
            "errors": result.get("errors", []),
        }
    except Exception as exc:
        raise HTTPException(500, str(exc))


@app.get("/api/intraday")
def api_intraday():
    snapshot = _safe(da.load_intraday_snapshot, {})
    decisions = _safe(da.load_daily_decisions, {})
    attempts = _safe(lambda: da.all_decision_attempts(decisions), [])
    from collections import Counter
    signals = Counter(str(a.get("signal", "NONE")).upper() for a in attempts)
    decision_map = decisions.get("decisions", {}) if isinstance(decisions, dict) else {}
    symbols = sorted(decision_map)
    scans = snapshot.get("scans", []) if isinstance(snapshot, dict) else []
    latest_scan = scans[-1] if scans else {}
    return {
        "metrics": {
            "date": str(decisions.get("date", "-")) if isinstance(decisions, dict) else "-",
            "symbols_evaluated": len(symbols),
            "attempts": len(attempts),
            "buys": signals["BUY"], "sells": signals["SELL"],
            "latest_signals": latest_scan.get("signals_detected", 0),
        },
        "attempts": attempts,
        "symbols": symbols,
    }


@app.get("/api/trades")
def api_trades():
    positions = _safe(da.load_tracked_positions, [])
    snapshot = _safe(da.load_intraday_snapshot, {})
    requests = _safe(da.load_order_requests, [])
    stock_tv = _safe(load_stock_tradingview_state, {})
    from src.execution import order_requests as oq
    executed = snapshot.get("executed", []) if isinstance(snapshot, dict) else []
    skipped  = snapshot.get("skipped", []) if isinstance(snapshot, dict) else []
    sections = _safe(lambda: da.latest_trade_log_sections(12), [])
    worker_alive = _safe(oq.worker_is_alive, False)
    worker_age = _safe(oq.worker_age_seconds, None)
    open_req = sum(1 for r in requests if str(r.get("status")) in {"pending", "processing"})
    return {
        "positions": positions,
        "executed": executed,
        "skipped": skipped,
        "order_requests": requests,
        "open_requests": open_req,
        "worker_alive": worker_alive,
        "worker_age": worker_age,
        "trade_log": sections,
        "stock_tradingview": {
            "updated_at": stock_tv.get("updated_at") if isinstance(stock_tv, dict) else None,
            "executions": (stock_tv.get("executions", []) if isinstance(stock_tv, dict) else [])[:20],
        },
    }


@app.get("/api/orders-journal")
def api_orders_journal(limit: int = 500):
    try:
        broker_sync = reconcile_ibkr_open_orders()
        payload = load_order_journal(limit=max(1, min(1000, int(limit or 500))))
        if isinstance(payload, dict):
            payload["broker_sync"] = broker_sync
        return payload
    except Exception as exc:
        raise HTTPException(500, str(exc))


@app.post("/api/orders-journal/export")
async def api_orders_journal_export(body: dict):
    """Export the rows currently visible in the Orders Journal as XLSX."""
    rows = body.get("rows") if isinstance(body, dict) else None
    if not isinstance(rows, list):
        raise HTTPException(400, "rows must be a list")
    if len(rows) > 1000:
        raise HTTPException(400, "A maximum of 1000 rows can be exported")

    columns = [
        ("Symbol", "symbol"), ("Market", "market"), ("Source", "source"),
        ("Side", "side"), ("Status", "status"), ("Entry Plan", "planned_entry"),
        ("Entry Fill", "actual_entry"), ("Current Stop", "current_stop"),
        ("Planned Stop", "planned_stop"), ("Current Target", "current_target"),
        ("Planned Target", "planned_target"), ("Exit", "actual_exit"),
        ("Quantity", "quantity"), ("P&L", "pnl"), ("R Multiple", "r_multiple"),
        ("Opened", "opened_at"), ("Closed", "closed_at"), ("Review", "review"),
        ("Exit Reason", "exit_reason"), ("Execution ID", "execution_id"),
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orders Journal"
    sheet.append([label for label, _ in columns])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="16343B")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for raw in rows:
        if not isinstance(raw, dict):
            continue
        values = []
        for _, key in columns:
            value = raw.get(key)
            if key == "review":
                review = value if isinstance(value, dict) else {}
                value = review.get("grade") or review.get("review_status") or "Unreviewed"
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False)
            values.append(value)
        sheet.append(values)
    for column in ("F", "G", "H", "I", "J", "K", "L", "N"):
        for cell in sheet[column][1:]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '$#,##0.0000'
    for cell in sheet["O"][1:]:
        if isinstance(cell.value, (int, float)):
            cell.number_format = '0.00R'
    for column_cells in sheet.columns:
        values = ["" if cell.value is None else str(cell.value) for cell in column_cells]
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max((len(value) for value in values), default=0) + 2, 12), 28)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"orders_journal_{datetime.now(ET).strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.post("/api/orders-journal/review")
async def api_orders_journal_review(body: dict):
    journal_id = str(body.get("id") or body.get("journal_id") or "").strip()
    if not journal_id:
        raise HTTPException(400, "id is required")
    try:
        review = save_review(journal_id, body)
        return {"ok": True, "id": journal_id, "review": review}
    except Exception as exc:
        raise HTTPException(500, str(exc))


@app.post("/api/order/place")
async def api_order_place(body: dict):
    from src.execution import order_requests as oq
    from src.config import SETTINGS
    if not body.get("symbol"):
        raise HTTPException(400, "symbol is required")
    try:
        request_id = oq.submit_place(body, live=False)
        return {"ok": True, "id": request_id, "dry_run": SETTINGS.dry_run_mode, "message": f"Queued place order for {body['symbol']} — paper only, respects DRY_RUN."}
    except Exception as exc:
        raise HTTPException(500, str(exc))


@app.post("/api/order/close")
async def api_order_close(body: dict):
    from src.execution import order_requests as oq
    from src.config import SETTINGS
    symbol = str(body.get("symbol", "")).upper()
    if not symbol:
        raise HTTPException(400, "symbol is required")
    try:
        request_id = oq.submit_close(symbol, live=False)
        return {"ok": True, "id": request_id, "dry_run": SETTINGS.dry_run_mode, "message": f"Queued close for {symbol} — paper only, respects DRY_RUN."}
    except Exception as exc:
        raise HTTPException(500, str(exc))


@app.get("/api/reports")
def api_reports():
    infos = _safe(da.list_report_info, [])
    return {"reports": [
        {"file": r.path.name, "kind": r.kind, "date": r.report_date,
         "updated_at": r.updated_at.isoformat(), "size": r.size_bytes}
        for r in infos
    ]}


@app.post("/api/forecast/run")
async def api_forecast_run(body: dict):
    try:
        scenario = fc.ForecastScenario(
            account_equity=float(body.get("equity", 100000)),
            cash_available=float(body.get("cash", 100000)),
            daily_realized_pnl=float(body.get("dailyPnl", 0)),
            assumed_spread_pct=float(body.get("spread", 0.05)) / 100,
            risk_per_trade=float(body.get("riskTrade", 1.0)) / 100,
            max_daily_loss_pct=float(body.get("maxDailyLoss", 2.0)) / 100,
            min_reward_risk_ratio=float(body.get("minRR", 3.0)),
            max_positions=int(body.get("maxPositions", 5)),
            max_spread_pct=float(body.get("maxSpread", 0.20)) / 100,
            max_open_risk_pct=float(body.get("maxOpenRisk", 6.0)) / 100,
            max_position_value=float(body.get("maxPosValue", 50000)),
            min_projected_entry_distance_atr=float(body.get("minEntryAtr", 1.0)),
            max_projected_entry_distance_atr=float(body.get("maxEntryAtr", 2.0)),
        )
        watchlist = _safe(da.load_watchlist, {})
        positions = _safe(da.load_tracked_positions, [])
        projected = _safe(lambda: fc.projected_level_candidates(watchlist), [])
        active, warns = _safe(lambda: fc.replay_active_candidates(watchlist, da.get_bars), ([], []))
        candidates = active + projected
        result = fc.run_forecast(candidates, scenario, positions)
        # make serialisable
        import math
        def _clean(v):
            if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
                return None
            return v
        def _fix(obj):
            if isinstance(obj, dict):  return {k: _fix(v) for k, v in obj.items()}
            if isinstance(obj, list):  return [_fix(i) for i in obj]
            return _clean(obj)
        return _fix(result)
    except Exception as exc:
        raise HTTPException(500, str(exc))


if __name__ == "__main__":
    uvicorn.run("dashboard_react.server:app", host="127.0.0.1", port=8550, reload=False)
